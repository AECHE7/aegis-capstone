import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_it_guide():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_inst.add_run('CENTRAL LUZON STATE UNIVERSITY\n')
    r1.font.name = 'Arial'
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(15, 61, 35)

    r2 = p_inst.add_run('College of Engineering • Department of Information Technology\nScience City of Muñoz, Nueva Ecija, Philippines\n')
    r2.font.name = 'Arial'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(80, 80, 80)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run('A.E.G.I.S. IT EXPERT TECHNICAL EVALUATION GUIDE\n')
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(13)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 61, 35)

    r_sub = p_title.add_run('System Architecture, Security Hardening & ISO/IEC 25010:2023 Assessment Protocol\nDocument Code: CLSU-CEn-DIT-AEGIS-GUIDE-IT-2026')
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(9)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph('_________________________________________________________________________________')

    # Content
    doc.add_heading('1. Executive Technical Overview', level=2)
    doc.add_paragraph(
        'A.E.G.I.S. is an enterprise scholarship governance platform featuring dual-pipeline AI document forensics, '
        'automated email notifications, and tamper-evident audit logging for the Central Luzon State University Office of Student Affairs (OSA).'
    )

    doc.add_heading('Technical Architecture & Specifications:', level=3)
    doc.add_paragraph(
        '• Web Tier: Laravel 12 (PHP 8.2), Blade Engine, Vanilla CSS & Tailwind CSS, Bootstrap 5.\n'
        '• Forensic AI Tier: Python Flask microservice integrating Error Level Analysis (ELA 95%), ResNet-50 CNN, SIFT clone detection, and Grad-CAM localized heatmaps.\n'
        '• Data & Storage Tier: SQLite (Dev) / PostgreSQL & MySQL 8.0 (Prod), Cloudflare R2 / S3 object storage, AES-256 database column encryption.\n'
        '• Statutory Controls: R.A. 10173 (Data Privacy Act of 2012) compliant, SHA-256 zero-knowledge OTP hashing at rest, immutable action logging.'
    )

    doc.add_heading('2. Evaluation Accounts Cheat Sheet', level=2)
    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Role', 'Email', 'Password', 'Security & MFA Behavior']
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F3D23"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ('SuperAdmin / Director', 'director@clsu.edu.ph', 'password', 'Auto-Bypassed (Demo Mode)'),
        ('Admin / Staff Evaluator', 'admin@clsu.edu.ph', 'password', 'Auto-Bypassed (Demo Mode)'),
        ('Student (Active History)', 'student@clsu.edu.ph', 'password', 'Demo OTP: 123456 or 000000'),
        ('Student (Clean /apply)', 'student_apply@clsu.edu.ph', 'password', 'Demo OTP: 123456 or 000000'),
    ]
    for r_idx, row in enumerate(data, start=1):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    doc.add_heading('3. Technical Verification Procedures (TC-IT-01 to TC-IT-10)', level=2)
    scenarios = [
        ('TC-IT-01: Authentication Hardening & SQLi Protection', 'Inject SQL syntax on /login (e.g. \' OR 1=1--). Verify Eloquent PDO parameter binding, bcrypt cost=12 hashing, and 5 attempts/min rate limiting returning HTTP 429.'),
        ('TC-IT-02: MFA & Zero-Knowledge Hashing', 'Trigger 6-digit OTP delivery. Verify database storage as 64-character SHA-256 hash at rest, 10-minute dynamic TTL, and test code 123456 acceptance.'),
        ('TC-IT-03: AES-256 Column Encryption at Rest', 'Inspect student_profiles table in DB. Confirm clsu_id_number, contact_number, and guardian info are encrypted with AES-256-CBC and decrypted in-memory only.'),
        ('TC-IT-04: RBAC & Authorization Gates', 'Authenticate as Student and attempt direct URL navigation to /admin/dashboard, /superadmin/scholarships, and /superadmin/settings. Verify HTTP 403 / unauthorized block.'),
        ('TC-IT-05: AI Forensic ELA-CNN Pipeline', 'Submit or inspect tampered Certificate of Grades. Verify Fraud Probability Score (0-100%), risk badge, ELA difference map, and Grad-CAM explainability heatmap overlay.'),
        ('TC-IT-06: Tamper-Evident Audit Logging', 'Perform status update or export. Navigate to /superadmin/audit-logs. Verify immutable structured capture of actor identity, IP, user-agent hash, timestamp, and JSON diffs.'),
        ('TC-IT-07: Transport & Session Security Headers', 'Inspect network headers in DevTools. Confirm HSTS, X-Frame-Options, X-Content-Type-Options, CSP, and cookies marked HttpOnly, Secure, SameSite=Lax.'),
        ('TC-IT-08: Isolated PDF & Print Engine', 'Trigger evaluation sheet print. Verify exact 1-page letter layout without blank pages, high-res seal, and cryptographic validation QR code leading to verification endpoint.'),
        ('TC-IT-09: Asynchronous Queuing & Fault Recovery', 'Dispatch batch AI analysis. Verify background queue execution, zero UI freeze, and exponential retry schedules ([15s, 45s, 90s, 180s, 360s]) for container cold starts.'),
        ('TC-IT-10: Database Portability & Asset Optimization', 'Verify cross-DB migration integrity (SQLite/MySQL/PostgreSQL), eager loading query efficiency (zero N+1 queries), and production Vite Gzip asset compression (>75%).'),
    ]
    for title, desc in scenarios:
        p = doc.add_paragraph()
        r = p.add_run(f'• {title}: ')
        r.font.bold = True
        p.add_run(desc)

    doc.add_heading('4. ISO/IEC 25010:2023 Product Quality Rating', level=2)
    doc.add_paragraph(
        'Please rate each of the 24 technical statements in Part 2 of the evaluation form using the 5-point Likert scale:\n'
        '5 - Strongly Agree (Exceeds expectations)\n'
        '4 - Agree (Meets all technical standards)\n'
        '3 - Neither Agree nor Disagree (Acceptable)\n'
        '2 - Disagree (Technical revision needed)\n'
        '1 - Strongly Disagree (Critical flaw detected)\n'
        'N/A - Not Applicable\n\n'
        'Complete Part 3 by indicating your overall acceptance recommendation and affixing your signature.'
    )

    doc.save('docs/AEGIS_IT_Expert_Evaluation_Instructional_Guide.docx')
    print('Successfully generated docs/AEGIS_IT_Expert_Evaluation_Instructional_Guide.docx')

if __name__ == '__main__':
    create_it_guide()
