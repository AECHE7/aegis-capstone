import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_shading(cell, color_hex):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{b}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        borders.append(border)
    tblPr.append(borders)

def style_heading(p, text, level=1):
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(16, 85, 45) # CLSU Green
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
    elif level == 3:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    return run

def create_document():
    doc = docx.Document()
    
    # Page setup - 0.7 inch margins for clean standard form spacing
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
    
    # Header Information
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_after = Pt(2)
    r1 = p_header.add_run("CENTRAL LUZON STATE UNIVERSITY\n")
    r1.font.name = 'Calibri'
    r1.font.size = Pt(13)
    r1.bold = True
    r1.font.color.rgb = RGBColor(16, 85, 45)
    
    r2 = p_header.add_run("College of Engineering • Department of Information Technology\n")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(10)
    r2.bold = True
    
    r3 = p_header.add_run("Science City of Muñoz, Nueva Ecija, Philippines\n")
    r3.font.name = 'Calibri'
    r3.font.size = Pt(9)
    r3.italic = True
    r3.font.color.rgb = RGBColor(100, 116, 139)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(6)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("―" * 55)
    r_div.font.color.rgb = RGBColor(203, 213, 225)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    rt = p_title.add_run("A.E.G.I.S. CAPSTONE COMPLIANCE:\nCOMPREHENSIVE SYSTEM TESTING AND ISO/IEC 25010:2023 EVALUATION SUITE\n")
    rt.font.name = 'Calibri'
    rt.font.size = Pt(12)
    rt.bold = True
    rt.font.color.rgb = RGBColor(15, 23, 42)
    
    r_sub = p_title.add_run("Standardized Quality Assessment Instruments for IT Experts, OSA Director, Staff Evaluators, and Student Applicants")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(9.5)
    r_sub.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, "CCCCCC")
    meta_data = [
        ("Project Title:", "A.E.G.I.S: AI-ENHANCED GRANT INFORMATION SYSTEM WITH DOCUMENT FORENSICS AND AUTOMATED NOTIFICATION FOR THE OFFICE OF STUDENT AFFAIRS"),
        ("Target Stakeholders:", "1. IT Experts & Faculty  |  2. OSA Director (SuperAdmin)  |  3. Staff Evaluators  |  4. Student Applicants"),
        ("Researchers & Leaders:", "John Andrei Carillo, Noriel S. Gadiano, Joshua A. Razon (BSIT - College of Engineering)"),
        ("Project Advisers:", "Louise Gwendolyn B. Hidalgo, MIT; Inigo Gabriel M. Balmadrid, MIT; Joseph Ariel J. Barza, MIT")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(5.3)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_shading(c0, "F1F5F9")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(8.5)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(8.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Regulatory Notice
    p_notice = doc.add_paragraph()
    p_notice.paragraph_format.space_after = Pt(8)
    r_not = p_notice.add_run("REGULATORY COMPLIANCE & PRIVACY NOTICE: ")
    r_not.font.name = 'Calibri'
    r_not.font.size = Pt(8.5)
    r_not.bold = True
    r_not_body = p_notice.add_run("This evaluation instrument conforms strictly to ISO/IEC 25010:2023 SQuaRE product quality standards, Republic Act No. 10173 (Data Privacy Act of 2012), and R.A. 11032 (Ease of Doing Business). Participation is voluntary. Gathered responses are protected under data minimization principles and used solely for academic research, system validation, and statistical analysis.")
    r_not_body.font.name = 'Calibri'
    r_not_body.font.size = Pt(8)
    r_not_body.font.color.rgb = RGBColor(100, 116, 139)

    # Helper function to add evaluator metadata row
    def add_evaluator_header(section_title, role_desc, fields):
        doc.add_page_break()
        h1 = doc.add_paragraph()
        style_heading(h1, section_title, level=1)
        
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(4)
        r_sub = p_sub.add_run(role_desc)
        r_sub.font.name = 'Calibri'
        r_sub.font.size = Pt(9)
        r_sub.italic = True
        r_sub.font.color.rgb = RGBColor(71, 85, 105)
        
        tbl = doc.add_table(rows=len(fields), cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, "CCCCCC")
        for i, (f1_k, f1_v, f2_k, f2_v) in enumerate(fields):
            row = tbl.rows[i]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(3.55)
            c1.width = Inches(3.55)
            set_cell_margins(c0, 50, 50, 80, 80)
            set_cell_margins(c1, 50, 50, 80, 80)
            
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0_k = p0.add_run(f1_k + ": ")
            r0_k.font.name = 'Calibri'
            r0_k.font.size = Pt(8.5)
            r0_k.bold = True
            r0_v = p0.add_run(f1_v)
            r0_v.font.name = 'Calibri'
            r0_v.font.size = Pt(8.5)
            
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1_k = p1.add_run(f2_k + ": ")
            r1_k.font.name = 'Calibri'
            r1_k.font.size = Pt(8.5)
            r1_k.bold = True
            r1_v = p1.add_run(f2_v)
            r1_v.font.name = 'Calibri'
            r1_v.font.size = Pt(8.5)
            
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Helper function to add test scenarios table
    def add_scenario_table(title, scenarios):
        p_t = doc.add_paragraph()
        style_heading(p_t, title, level=2)
        
        tbl = doc.add_table(rows=len(scenarios) + 1, cols=6)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, "B0B0B0")
        
        headers = ["No.", "Feature / Module", "Test Scenario & Protocol", "Expected Result", "Status", "Remarks"]
        col_widths = [Inches(0.4), Inches(1.3), Inches(2.2), Inches(1.7), Inches(0.8), Inches(0.7)]
        
        # Header Row
        hdr_row = tbl.rows[0]
        for c_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_shading(cell, "0F3D23") # Dark green
            set_cell_margins(cell, 80, 80, 60, 60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            r.font.name = 'Calibri'
            r.font.size = Pt(8)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
        for r_idx, (num, feat, scen, exp) in enumerate(scenarios, start=1):
            row = tbl.rows[r_idx]
            for c_idx, cell in enumerate(row.cells):
                cell.width = col_widths[c_idx]
                set_cell_margins(cell, 50, 50, 50, 50)
                if r_idx % 2 == 0:
                    set_cell_shading(cell, "F8FAFC")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                
            row.cells[0].paragraphs[0].add_run(str(num)).font.size = Pt(8)
            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            r_feat = row.cells[1].paragraphs[0].add_run(feat)
            r_feat.font.name = 'Calibri'
            r_feat.font.size = Pt(8)
            r_feat.bold = True
            
            r_scen = row.cells[2].paragraphs[0].add_run(scen)
            r_scen.font.name = 'Calibri'
            r_scen.font.size = Pt(7.5)
            
            r_exp = row.cells[3].paragraphs[0].add_run(exp)
            r_exp.font.name = 'Calibri'
            r_exp.font.size = Pt(7.5)
            
            r_stat = row.cells[4].paragraphs[0].add_run("☐ Pass\n☐ Fail\n☐ Rev.\n☐ N/A")
            r_stat.font.name = 'Calibri'
            r_stat.font.size = Pt(7)
            
            row.cells[5].paragraphs[0].add_run("").font.size = Pt(7)
            
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Helper function to add ISO 25010 Questionnaire table
    def add_iso_table(title, subtitle, statements):
        p_t = doc.add_paragraph()
        style_heading(p_t, title, level=2)
        
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_after = Pt(4)
        r_s = p_s.add_run(subtitle + "\nRating Scale: 5 - Strongly Agree (SA) | 4 - Agree (A) | 3 - Neutral (N) | 2 - Disagree (D) | 1 - Strongly Disagree (SD) | N/A - Not Applicable")
        r_s.font.name = 'Calibri'
        r_s.font.size = Pt(8)
        r_s.italic = True
        r_s.font.color.rgb = RGBColor(100, 116, 139)
        
        tbl = doc.add_table(rows=len(statements) + 1, cols=3)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, "B0B0B0")
        
        col_widths = [Inches(0.5), Inches(5.1), Inches(1.5)]
        hdr_row = tbl.rows[0]
        for c_idx, h_text in enumerate(["No.", "ISO/IEC 25010:2023 Evaluation Statement", "Rating"]):
            cell = hdr_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_shading(cell, "1E293B") # Dark slate
            set_cell_margins(cell, 70, 70, 60, 60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
        for r_idx, item in enumerate(statements, start=1):
            row = tbl.rows[r_idx]
            c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
            c0.width, c1.width, c2.width = col_widths[0], col_widths[1], col_widths[2]
            set_cell_margins(c0, 45, 45, 50, 50)
            set_cell_margins(c1, 45, 45, 60, 60)
            set_cell_margins(c2, 45, 45, 50, 50)
            
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_after = Pt(0)
            
            if len(item) == 2: # Category Header
                cat_num, cat_title = item
                set_cell_shading(c0, "E2E8F0")
                set_cell_shading(c1, "E2E8F0")
                set_cell_shading(c2, "E2E8F0")
                r0 = p0.add_run(cat_num)
                r0.font.name = 'Calibri'
                r0.font.size = Pt(8)
                r0.bold = True
                r1 = p1.add_run(cat_title)
                r1.font.name = 'Calibri'
                r1.font.size = Pt(8.5)
                r1.bold = True
                p2.add_run("")
            else: # Statement Row
                stmt_num, stmt_text = item[0], item[1]
                if r_idx % 2 == 0:
                    set_cell_shading(c0, "F8FAFC")
                    set_cell_shading(c1, "F8FAFC")
                    set_cell_shading(c2, "F8FAFC")
                r0 = p0.add_run(stmt_num)
                r0.font.name = 'Calibri'
                r0.font.size = Pt(8)
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                r1 = p1.add_run(stmt_text)
                r1.font.name = 'Calibri'
                r1.font.size = Pt(8)
                
                r2 = p2.add_run("☐ 5  ☐ 4  ☐ 3  ☐ 2  ☐ 1  ☐ N/A")
                r2.font.name = 'Calibri'
                r2.font.size = Pt(7)
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Helper function to add sign-off block
    def add_signoff(eval_title, is_director=False):
        p_res = doc.add_paragraph()
        style_heading(p_res, "Evaluation Result & Sign-Off Endorsement", level=2)
        
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_after = Pt(4)
        if is_director:
            p_opt.add_run("Overall System Acceptance Decision:\n☐ ACCEPTED FOR LIVE INSTITUTIONAL USE — meets all executive and operational requirements.\n☐ ACCEPTED WITH MINOR REVISIONS — usable subject to noted enhancements.\n☐ FOR REVISION AND RETESTING — critical governance/functional issues must be resolved before deployment.")
        else:
            p_opt.add_run("Overall System Acceptance Decision:\n☐ ACCEPTED — operates satisfactorily; no critical issue prevents intended operational use.\n☐ ACCEPTED WITH MINOR REVISIONS — system is usable, subject to noted engineering/operational adjustments.\n☐ FOR REVISION AND RETESTING — defects or usability hurdles must be corrected prior to acceptance.")
        p_opt.runs[0].font.name = 'Calibri'
        p_opt.runs[0].font.size = Pt(8)
        
        p_com = doc.add_paragraph()
        p_com.paragraph_format.space_after = Pt(4)
        r_c = p_com.add_run("Evaluator Comments & Recommendations:\n________________________________________________________________________________________________________________________\n________________________________________________________________________________________________________________________")
        r_c.font.name = 'Calibri'
        r_c.font.size = Pt(8)
        
        tbl_sig = doc.add_table(rows=2, cols=2)
        tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_sig, "CCCCCC")
        
        c0, c1 = tbl_sig.rows[0].cells[0], tbl_sig.rows[0].cells[1]
        c0.width, c1.width = Inches(3.55), Inches(3.55)
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run("_____________________________________________________\n").font.size = Pt(8.5)
        r0 = p0.add_run(f"Evaluator Signature over Printed Name\n({eval_title})")
        r0.font.name = 'Calibri'
        r0.font.size = Pt(8)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run("_____________________________________________________\n").font.size = Pt(8.5)
        r1 = p1.add_run("Project Leaders / Student Researchers\n(Carillo, Gadiano, Razon)")
        r1.font.name = 'Calibri'
        r1.font.size = Pt(8)
        r1.bold = True
        
        c2, c3 = tbl_sig.rows[1].cells[0], tbl_sig.rows[1].cells[1]
        c2.width, c3.width = Inches(3.55), Inches(3.55)
        set_cell_margins(c2, 40, 40, 80, 80)
        set_cell_margins(c3, 40, 40, 80, 80)
        
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run("Date: ________________________").font.size = Pt(8)
        
        p3 = c3.paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.add_run("Date: ________________________").font.size = Pt(8)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # SECTION 1: IT EXPERT EVALUATION
    # ==========================================
    add_evaluator_header(
        "SECTION 1: IT EXPERT TESTING & ISO/IEC 25010 EVALUATION INSTRUMENT",
        "For Technical Validation by Software Architects, Cybersecurity Auditors, AI/ML Engineers, and IT Faculty",
        [
            ("Evaluator Name", "________________________________________", "Position / Designation", "________________________"),
            ("Organization / Agency", "________________________________________", "Specialization", "☐ Arch  ☐ Sec  ☐ AI  ☐ Faculty"),
            ("Date of Testing", "October 2026", "Environment / Build", "Cloud Staging Portal (v1.0.0)")
        ]
    )
    
    it_scenarios = [
        (1, "Authentication & Bcrypt Hardening", "Inject SQL payloads (' OR 1=1--) and brute-force /login. Verify Bcrypt (cost=12), rate-limiter (5/min), and CSRF tokens.", "SQL injection blocked; password hashes stored with irreversible Bcrypt; IP rate-limiter returns HTTP 429 upon threshold breach."),
        (2, "MFA & Zero-Knowledge Hashing", "Trigger 6-digit OTP delivery. Inspect db column otp_code. Verify SHA-256 hash storage at rest and 10-min countdown timer.", "OTP stored strictly as 64-character SHA-256 hash; expired tokens rejected; universal demo OTP accepted for test fixtures."),
        (3, "AES-256 Column Encryption", "Inspect raw db storage for sensitive student profile attributes (CLSU ID, guardian contacts) under R.A. 10173 data minimization.", "Sensitive data stored as AES-256-CBC ciphertext; raw SQL returns ciphertext; in-memory decryption executed only for authorized sessions."),
        (4, "RBAC & Authorization Gates", "Authenticate as Student and attempt direct URL navigation to privileged routes (/admin/dashboard, /superadmin/settings).", "Unauthorized navigation intercepted by CheckRole middleware; returns HTTP 403 Forbidden; zero data leakage."),
        (5, "AI ELA-CNN Forensic Pipeline", "Submit tampered COG with altered numerical grades. Execute ELA (Q=95), ResNet-50 inference, and inspect Grad-CAM heatmap.", "Pipeline computes ELA difference map, returns Fraud Probability Score (0-100%), maps risk tier, and highlights edited grade bounding box."),
        (6, "Tamper-Evident Audit Logging", "Execute administrative operations (approve application, edit settings, export data). Verify structured audit trail records.", "Immutable audit logs created in admin_action_logs and export_access_logs with actor ID, IP address, timestamp, and JSON diffs."),
        (7, "Session & Header Hardening", "Inspect HTTP response headers and cookie flags via DevTools on authenticated HTTPS connection.", "HSTS, X-Frame-Options: SAMEORIGIN, nosniff, and CSP active; session cookies flagged Secure, HttpOnly, SameSite=Lax."),
        (8, "PDF & Isolated Print Engine", "Generate 1-page applicant evaluation form with verification checklist, QR badge, and seal. Trigger isolated iframe print engine.", "High-fidelity 1-page letter PDF generated matching web modal preview 1:1; isolated iframe prints cleanly without blank pages; QR valid."),
        (9, "Queue Worker & Fault Recovery", "Trigger batch AI forensic scans. Inspect database queue worker, exponential retry backoff, and cold-start tolerance.", "Asynchronous worker processes jobs without blocking web UI; container cold starts recover gracefully via retry backoff."),
        (10, "Portability & Asset Optimization", "Verify DB portability across SQLite/PostgreSQL. Inspect eager loading (no N+1 queries) and Gzip compression on assets.", "Schema executes cleanly on both databases; zero N+1 queries during application indexing; Gzip reduces assets by >75%.")
    ]
    add_scenario_table("Part 1.A: Technical Testing & Verification Matrix", it_scenarios)
    
    it_statements = [
        ("1", "FUNCTIONAL SUITABILITY"),
        ("1.1", "The system completely implements all required scholarship modules, AI forensic inspection, and notification lifecycles."),
        ("1.2", "The system executes calculation algorithms (GWA checks, quota thresholds, AI risk scoring) with mathematical precision."),
        ("1.3", "Functional workflows provide appropriate administrative utilities (bulk actions, audit trails) without redundant steps."),
        ("2", "PERFORMANCE EFFICIENCY"),
        ("2.1", "The system responds within acceptable latency thresholds (< 1.5 seconds) during typical database queries and page transitions."),
        ("2.2", "Server resources (CPU, RAM, database I/O) are utilized efficiently via eager loading, query caching, and asset compression."),
        ("2.3", "The system manages concurrent applicant uploads and background worker queuing without deadlocks or performance degradation."),
        ("3", "COMPATIBILITY"),
        ("3.1", "The application co-exists smoothly in multi-container cloud environments (Docker, Alpine Linux, Nginx) without resource contention."),
        ("3.2", "The system integrates seamlessly with external web services and APIs (Brevo SMTP, Cloudflare R2, Python AI microservices)."),
        ("4", "USABILITY (INTERACTION CAPABILITY)"),
        ("4.1", "The software provides clear architectural recognizability, intuitive UI patterns, breadcrumbs, and standardized SweetAlert2 dialogs."),
        ("4.2", "The interfaces enable rapid user learnability with minimal training through structured multi-step wizards and guided tours."),
        ("4.3", "The application enforces robust client-side and server-side validation rules with accessible error handling to prevent user mistakes."),
        ("4.4", "The user interface is cleanly styled with modern typography, responsive viewports, and high contrast adhering to accessibility standards."),
        ("5", "RELIABILITY"),
        ("5.1", "The software demonstrates architectural maturity, passing automated test suites with high code coverage and zero regressions."),
        ("5.2", "The system provides high operational availability with fault-tolerant fallbacks (database connection retries, email failover logging)."),
        ("5.3", "The system gracefully recovers from service interruptions (e.g. AI container cold starts) through automated job retry backoffs."),
        ("6", "SECURITY"),
        ("6.1", "Unauthorized access to sensitive student records is strictly prevented through role-based middleware and AES-256 database encryption."),
        ("6.2", "Data integrity is robustly safeguarded through CSRF token verification, cryptographic URL signatures, and immutable audit logs."),
        ("6.3", "System activities (approvals, rejections, setting updates, exports) are immutably tied to user identity for complete non-repudiation."),
        ("6.4", "User identity is verified through cryptographically secure Multi-Factor Authentication (MFA OTP) with SHA-256 hash storage at rest."),
        ("7", "MAINTAINABILITY"),
        ("7.1", "The codebase exhibits high modularity following MVC and Clean Architecture standards (Skinny Controllers, Fat Models, Service Layer)."),
        ("7.2", "Software components (forensic analyzers, notification dispatchers, PDF generators, alert engines) are abstracted for reusability."),
        ("7.3", "The codebase provides clear architectural documentation, structured logging, and high testability with automated test coverage."),
        ("8", "PORTABILITY"),
        ("8.1", "The web application adapts seamlessly across modern web browsers (Chrome, Edge, Safari, Firefox) and multi-device form factors."),
        ("8.2", "The application deployment pipeline is standardized via containerization (Dockerfile, render.yaml) and automated database migrations.")
    ]
    add_iso_table("Part 1.B: ISO/IEC 25010:2023 Technical Product Quality Evaluation", "Evaluation Instrument — IT Experts & Technical Specialists", it_statements)
    add_signoff("IT Expert / Technical Specialist")

    # ==========================================
    # SECTION 2: OSA DIRECTOR EVALUATION
    # ==========================================
    add_evaluator_header(
        "SECTION 2: OSA DIRECTOR (SUPER ADMINISTRATOR) TESTING & EVALUATION INSTRUMENT",
        "Executive Governance, Policy Administration, Statutory Auditing & Institutional Sustainability",
        [
            ("Evaluator Name", "________________________________________", "Position / Designation", "Director / Head, Office of Student Affairs"),
            ("Client Agency", "Central Luzon State University — OSA", "Date of Evaluation", "October 2026"),
            ("Testing Portal", "Executive Director Dashboard (/superadmin/dashboard)", "Software Build", "Production Web Portal (v1.0.0)")
        ]
    )
    
    dir_scenarios = [
        (1, "Executive Analytics & Overhaul", "Navigate to /superadmin/analytics. Test 1-click active term filtering, slot quota capacity progress bars, GWA integrity index, and anomaly distribution.", "Metrics render dynamically; slot quotas calculate against real caps; anomaly keys display readable titles; active term filter works cleanly."),
        (2, "Academic Term & Semester Governance", "Navigate to /superadmin/settings. Review active academic terms. Create new semester, switch active term, and inspect topbar badge synchronization.", "System strictly enforces single active term rule; active term badge in topbar updates in real time; past terms archived safely."),
        (3, "Program Lifecycle Governance", "Create new scholarship program, configure eligibility thresholds (minimum GWA, eligible colleges, maximum slots), and design custom application fields.", "Program published instantly in catalog; quota enforced; custom form questions display correctly in student application form."),
        (4, "Staff Account & Queue Delegation", "Issue secure email invitation token for a new OSA scholarship evaluator. Assign specific grant programs to evaluator.", "Token-based activation link dispatched; invited evaluator registers securely; assigned queues route matching applications to evaluator."),
        (5, "Unified Communications Center", "Open Communications Center (/admin/announcements). Publish portal announcement, compose email broadcast with audience filters, inspect delivery logs.", "Portal announcement renders on student notice board; targeted email broadcast dispatches to filtered student cohorts; logs record delivery count."),
        (6, "Statutory Compliance Export Hub", "Navigate to Compliance Export Hub. Select date presets (This Year, Last 30d) and generate official CHED StuFAPs and DOST-SEI compliance reports in CSV and PDF.", "Formatted masterlist exports generated with applicant demographics, approved grant amounts, GWA ratings, and verification clearance timestamps."),
        (7, "System Settings & Calibration", "Access /superadmin/settings. Adjust AI Fraud Detection Threshold (e.g. 70%), GWA discrepancy tolerance (0.01), and MFA enforcement level.", "System dynamically updates global settings table; newly processed applications evaluate against new calibration values immediately."),
        (8, "System Trash & Soft-Deletion Recovery", "Soft-delete a test scholarship program. Navigate to System Trash /superadmin/trash. Verify program recovery and permanent purge controls.", "Soft-deleted entity hidden from public views; restored cleanly from Trash tab without data loss; audit log captures recovery event.")
    ]
    add_scenario_table("Part 2.A: Executive Governance & Compliance Testing Matrix", dir_scenarios)
    
    dir_statements = [
        ("1", "FUNCTIONAL SUITABILITY"),
        ("2.1", "The system provides complete executive tools for managing university scholarship programs, evaluator staffing, and semester terms."),
        ("2.2", "Real-time analytics, quota utilization tracking, and fraud indices produce accurate and dependable management information."),
        ("2.3", "The system effectively automates administrative workloads that previously required tedious manual physical record collation."),
        ("2", "PERFORMANCE EFFICIENCY"),
        ("2.4", "Executive dashboards and analytics visualizations load swiftly without perceptible delay during operational hours."),
        ("2.5", "Generation of extensive statutory compliance reports (CSV and PDF masterlists) executes rapidly."),
        ("3", "USABILITY (INTERACTION CAPABILITY)"),
        ("2.6", "Executive navigation menus, metric summary cards, and quick filter pills are intuitive and straightforward to navigate."),
        ("2.7", "The guided tour and onboarding cues effectively orient new administrators to portal capabilities."),
        ("2.8", "Information presentation is visually polished, well-organized, and legible across desktop and laptop screens."),
        ("4", "RELIABILITY"),
        ("2.9", "The system performs reliably without unexpected crashes, server errors, or interrupted operations."),
        ("2.10", "Automated email broadcasts and applicant notifications deliver consistently without dropped messages."),
        ("5", "SECURITY & NON-REPUDIATION"),
        ("2.11", "Access to confidential student financial and academic records is strictly safeguarded against unauthorized staff or external access."),
        ("2.12", "Comprehensive 7-tier audit logs maintain complete accountability by recording every approval, rejection, and configuration update."),
        ("2.13", "Multi-Factor Authentication (MFA) and trusted device controls offer dependable protection for administrative accounts."),
        ("6", "FLEXIBILITY & COMPLIANCE"),
        ("2.14", "The system readily accommodates changing institutional scholarship rules, new grant programs, and unique application forms."),
        ("2.15", "The system supports multi-semester management, allowing smooth transition between academic terms without data loss."),
        ("2.16", "Exported reports conform to CHED StuFAPs and DOST-SEI institutional auditing and reporting standards."),
        ("7", "SAFETY & RISK REDUCTION"),
        ("2.17", "The system provides clear confirmation dialogs before high-impact administrative actions (e.g. revoking grants, deleting terms)."),
        ("2.18", "The System Trash repository safeguards against accidental permanent deletion of valuable scholarship programs and records.")
    ]
    add_iso_table("Part 2.B: ISO/IEC 25010:2023 Executive Quality Evaluation", "Evaluation Instrument — Office of Student Affairs Director", dir_statements)
    add_signoff("Office of Student Affairs Director", is_director=True)

    # ==========================================
    # SECTION 3: OSA STAFF EVALUATION
    # ==========================================
    add_evaluator_header(
        "SECTION 3: OSA STAFF (SCHOLARSHIP EVALUATOR) TESTING & EVALUATION INSTRUMENT",
        "Operational Intake, Document Verification, AI Decision-Support & Award Processing",
        [
            ("Evaluator Name", "________________________________________", "Position / Designation", "OSA Scholarship Evaluator / Staff Officer"),
            ("Unit / Section", "Student Welfare & Scholarship Division", "Date of Evaluation", "October 2026"),
            ("Testing Module", "Evaluator Application Queue (/admin/applications)", "Software Build", "Production Web Portal (v1.0.0)")
        ]
    )
    
    staff_scenarios = [
        (1, "Application Triage & Queues", "Access /admin/applications. Filter queue by program, status (Under Review, Returned), college, and search by student ID number.", "Queue instantaneously filters matching submissions; priority badges display clearly; after-hours submission flags visible."),
        (2, "Review Canvas & Dual Viewer", "Open applicant review canvas (/admin/review/{id}). Inspect side-by-side document inspection pane with zoom, pan, and contrast inversion controls.", "High-resolution preview of Certificate of Grades (COG) and Certificate of Registration (COR) renders smoothly with fluid zoom and pan."),
        (3, "OCR Grade Parsing & GWA Check", "Inspect OCR extracted grade sheet table. Verify parsed course codes, credit units, numerical grades, and computed GWA against student declared GWA.", "Parsed grade table highlights any mathematical discrepancies exceeding tolerance (0.01) with distinct warning badges."),
        (4, "AI Forensics & Heatmap Inspection", "Inspect AI Analysis Dossier on the review canvas. Review Fraud Probability Score (FPS), Risk Tier badge, and toggle Grad-CAM heatmap overlay.", "Color-coded risk tier displays prominently; Grad-CAM heatmap highlights suspicious altered pixel clusters; evaluator retains override authority."),
        (5, "Application Decisioning", "Process application decisions: (a) Approve grant, (b) Reject with standard justification, or (c) Return for Correction with specific deficiency remarks.", "Decision modal requires mandatory justification for rejections/returns; updates status immediately in database; deducts slot quota upon approval."),
        (6, "Automated Notification Dispatch", "Render an application decision. Verify applicant receives automated status email and in-app bell notification drawer update.", "Immediate automated email dispatched via Brevo SMTP containing official status, remarks, and next steps; bell drawer updates badge count."),
        (7, "Official 1-Page Evaluation PDF", "Click 'Generate Official Evaluation Form'. Inspect modal preview and trigger print engine via isolated iframe.", "Form renders in standard 1-page letter format with CLSU OSA seal, applicant credentials, checklist, QR verification badge, and signature lines."),
        (8, "Deficiency Resubmission", "Open an application resubmitted by a student following a 'Returned for Correction' status. Verify updated files and audit history.", "History drawer displays revision timeline; newly uploaded documents replace deficient files; evaluator can review corrections and approve.")
    ]
    add_scenario_table("Part 3.A: Operational Intake & AI Decision-Support Testing Matrix", staff_scenarios)
    
    staff_statements = [
        ("1", "FUNCTIONAL SUITABILITY"),
        ("3.1", "The review canvas provides all tools necessary to evaluate student credentials, inspect documents, and render decisions."),
        ("3.2", "OCR grade extraction and GWA discrepancy checks accurately detect mismatches between student input and official documents."),
        ("3.3", "The triage filtering tools enable efficient organization of applications by scholarship program and processing status."),
        ("2", "PERFORMANCE EFFICIENCY"),
        ("3.4", "Uploaded document images and PDF previews open swiftly without delaying the evaluation workflow."),
        ("3.5", "Recording an evaluation decision and updating the application queue occurs instantly upon submission."),
        ("3", "USABILITY (INTERACTION CAPABILITY)"),
        ("3.6", "The layout of the evaluation screen is logical, well-structured, and comfortable for extended daily use."),
        ("3.7", "Document zoom, pan, and rotation controls are easy and intuitive to operate."),
        ("3.8", "The interface effectively prevents accidental decisions by requiring explicit confirmation and remarks."),
        ("4", "AI FORENSIC UTILITY & EXPLAINABILITY"),
        ("3.9", "The Fraud Probability Score (0–100%) and color-coded risk tier badges provide clear, actionable guidance during review."),
        ("3.10", "The Grad-CAM explainability heatmap overlay helps locate specific suspicious areas on modified documents."),
        ("3.11", "The system appropriately positions AI as a decision-support aid, ensuring that the human evaluator retains full authority."),
        ("5", "RELIABILITY & ERROR RECOVERY"),
        ("3.12", "The application queue operates consistently without losing draft notes or entered evaluator comments."),
        ("3.13", "Automated notification emails trigger reliably whenever an application status changes."),
        ("6", "SECURITY & DATA PRIVACY"),
        ("3.14", "Student personal contact numbers, guardian records, and academic files are shielded from unauthorized exposure."),
        ("3.15", "Every evaluator action is logged with an immutable timestamp, ensuring fairness and accountability."),
        ("7", "SAFETY"),
        ("3.16", "The system requires mandatory correction remarks before returning an application, preventing student confusion."),
        ("3.17", "The system issues clear warnings before permanent rejection of an applicant's scholarship grant.")
    ]
    add_iso_table("Part 3.B: ISO/IEC 25010:2023 Operational Usability & AI Utility Evaluation", "Evaluation Instrument — OSA Scholarship Evaluator", staff_statements)
    add_signoff("OSA Scholarship Evaluator / Staff Officer")

    # ==========================================
    # SECTION 4: STUDENT APPLICANT EVALUATION
    # ==========================================
    add_evaluator_header(
        "SECTION 4: STUDENT APPLICANT (END-USER) TESTING & EVALUATION INSTRUMENT",
        "Student End-User Experience, Paperless Submissions, Mobile Responsiveness & Notifications",
        [
            ("Student Name (Optional)", "________________________________________", "College / Program", "________________________"),
            ("Year Level", "☐ 1st  ☐ 2nd  ☐ 3rd  ☐ 4th / Higher", "Device Used for Testing", "☐ Laptop/PC  ☐ Mobile  ☐ Tablet"),
            ("Date of Testing", "October 2026", "Testing Portal", "Student Scholarship Portal (/student)")
        ]
    )
    
    student_scenarios = [
        (1, "Registration & Profile Setup", "Register account using institutional @clsu2.edu.ph email, set a secure password, and complete academic profile.", "System enforces institutional email format, securely validates profile details, and grants access to student scholarship dashboard."),
        (2, "Catalog Discovery & Filtering", "Browse available scholarships on /scholarships. Filter by College, Year Level, and minimum GWA requirements.", "Catalog dynamically shows eligible grants, application deadlines, stipend amounts, and required documents."),
        (3, "Multi-Step Grant Application", "Select an open scholarship, complete program-specific questions, and review requirements checklist.", "Step-by-step application form guides user smoothly; indicates progress and prevents skipping required questions."),
        (4, "Document Upload & Validation", "Upload Certificate of Grades (COG) and Certificate of Registration (COR). Test file format validation and size limits.", "Clear upload drag-and-drop zone; client-side validator rejects oversized or invalid file extensions with helpful instructions."),
        (5, "Data Privacy Consent & Submit", "Review submission summary, accept Data Privacy Act (R.A. 10173) agreement, and submit application.", "Submission confirmed with instant success modal; status updates to 'Pending'; email confirmation received immediately."),
        (6, "Live Tracking & Announcements", "Track application status on Student Dashboard /student/dashboard. Check in-app notification drawer and announcements feed.", "Live progress tracker visually illustrates current stage (Pending -> Under Review -> Approved); announcements bulletin accessible."),
        (7, "Deficiency Correction", "Open an application marked 'Returned for Correction'. Read evaluator deficiency remarks, re-upload corrected COG, and resubmit.", "Deficiency instructions clearly explained; allows replacement of flagged documents without restarting entire application."),
        (8, "Award Clearance Download", "View notification for an approved scholarship. Preview and download official 1-page Scholarship Award & Clearance Certificate.", "High-quality 1-page certificate downloads with official university seal and QR code verification badge.")
    ]
    add_scenario_table("Part 4.A: Student End-User Workflow Testing Matrix", student_scenarios)
    
    student_statements = [
        ("1", "FUNCTIONAL SUITABILITY"),
        ("4.1", "The portal provides all functions needed to find, apply for, and monitor university scholarships."),
        ("4.2", "The system accurately calculates my submitted GWA and reflects my academic eligibility."),
        ("4.3", "The online process eliminates the need to submit physical paper folders at the university office."),
        ("2", "PERFORMANCE EFFICIENCY"),
        ("4.4", "Application pages, catalogs, and status trackers load quickly on my device."),
        ("4.5", "Uploading documents (grade sheets, registration cards) completes without long waiting times."),
        ("3", "COMPATIBILITY & MOBILE RESPONSIVENESS"),
        ("4.6", "The system functions properly on my smartphone, tablet, or laptop browser."),
        ("4.7", "The text, buttons, and upload forms remain readable and easy to tap on smaller mobile screens."),
        ("4", "INTERACTION CAPABILITY (USABILITY)"),
        ("4.8", "It is easy to understand how to use the portal without needing extensive training or instructions."),
        ("4.9", "Form instructions, requirements checklists, and deadline notices are clear and unambiguous."),
        ("4.10", "The system alerts me helpfully if I miss required questions or upload the wrong file format."),
        ("4.11", "The overall design, colors, and layout look modern, professional, and visually comfortable."),
        ("5", "RELIABILITY"),
        ("4.12", "The portal worked smoothly during submission without freezing or crashing."),
        ("4.13", "The system accurately preserves my entered application information and documents once submitted."),
        ("6", "SECURITY & DATA PRIVACY"),
        ("4.14", "I feel confident that my personal information, grades, and documents are securely protected."),
        ("4.15", "The system clearly informed me of my Data Privacy Act (R.A. 10173) rights prior to submission."),
        ("7", "COMMUNICATION & TRANSPARENCY"),
        ("4.16", "Email notifications and dashboard alerts kept me promptly updated regarding my application progress."),
        ("4.17", "If an application is returned for correction, the remarks clearly explain what documents must be updated."),
        ("8", "OVERALL SATISFACTION"),
        ("4.18", "Overall, I am highly satisfied with my experience using the A.E.G.I.S. scholarship portal."),
        ("4.19", "I would strongly prefer applying through this digital portal over traditional paper-based submissions."),
        ("4.20", "The system is ready for official university-wide implementation for all CLSU students.")
    ]
    add_iso_table("Part 4.B: ISO/IEC 25010:2023 End-User Usability & Quality Evaluation", "Evaluation Instrument — Student Applicants", student_statements)
    add_signoff("Student Applicant")

    # ==========================================
    # SECTION 5: SCORING GUIDE & STATISTICAL INTERPRETATION
    # ==========================================
    doc.add_page_break()
    p_s5 = doc.add_paragraph()
    style_heading(p_s5, "SECTION 5: STATISTICAL SCORING GUIDE & INTERPRETATION BENCHMARKS", level=1)
    
    p_guide_desc = doc.add_paragraph()
    p_guide_desc.paragraph_format.space_after = Pt(4)
    r_gd = p_guide_desc.add_run("For statistical analysis and manuscript reporting in Chapters 3, 4, and 5 of the Capstone Thesis, the following 5-point Likert scale interpretation benchmarks established by statistical conventions and ISO/IEC 25010 SQuaRE standards shall be applied:")
    r_gd.font.name = 'Calibri'
    r_gd.font.size = Pt(8.5)
    
    tbl_bench = doc.add_table(rows=6, cols=4)
    tbl_bench.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_bench, "B0B0B0")
    
    b_col_widths = [Inches(1.1), Inches(1.8), Inches(1.6), Inches(2.6)]
    b_headers = ["Mean Range", "Verbal Interpretation (Technical / Quality)", "Verbal Interpretation (Acceptability / End-User)", "Qualitative Standard Description"]
    
    for c_idx, h_text in enumerate(b_headers):
        cell = tbl_bench.rows[0].cells[c_idx]
        cell.width = b_col_widths[c_idx]
        set_cell_shading(cell, "0F3D23")
        set_cell_margins(cell, 70, 70, 60, 60)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = 'Calibri'
        r.font.size = Pt(8)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    bench_data = [
        ("4.21 – 5.00", "Strongly Agree / Exemplary Quality", "Very High Acceptability", "Quality, architectural controls, and UX significantly exceed operational requirements; flawless implementation."),
        ("3.41 – 4.20", "Agree / High Quality", "High Acceptability", "Satisfies all primary requirements and industry benchmarks with strong reliability and user satisfaction."),
        ("2.61 – 3.40", "Moderate Quality", "Moderate Acceptability", "System is functional and acceptable; minor architectural or cosmetic enhancements recommended."),
        ("1.81 – 2.60", "Disagree / Low Quality", "Low Acceptability", "Notable technical deficiencies or usability hurdles exist; corrective remediation required."),
        ("1.00 – 1.80", "Strongly Disagree / Very Low Quality", "Very Low Acceptability / Rejected", "Critical flaws, security breaches, or operational failures; system rejected in current form.")
    ]
    
    for r_idx, (m_rng, v_tech, v_user, q_desc) in enumerate(bench_data, start=1):
        row = tbl_bench.rows[r_idx]
        for c_idx in range(4):
            cell = row.cells[c_idx]
            cell.width = b_col_widths[c_idx]
            set_cell_margins(cell, 45, 45, 50, 50)
            if r_idx % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
                
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(m_rng)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(8)
        r0.bold = True
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v_tech)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(8)
        r1.bold = True
        
        p2 = row.cells[2].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(v_user)
        r2.font.name = 'Calibri'
        r2.font.size = Pt(8)
        
        p3 = row.cells[3].paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(q_desc)
        r3.font.name = 'Calibri'
        r3.font.size = Pt(7.5)
        
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(14)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("END OF COMPREHENSIVE TESTING AND EVALUATION SUITE\nDocument Code: CLSU-CEn-DIT-AEGIS-EVAL-2026-V1")
    r_end.font.name = 'Calibri'
    r_end.font.size = Pt(8)
    r_end.bold = True
    r_end.font.color.rgb = RGBColor(100, 116, 139)

    out_path = r"f:\aegis-capstone\docs\AEGIS_COMPREHENSIVE_TESTING_AND_EVALUATION_FORMS.docx"
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_document()
