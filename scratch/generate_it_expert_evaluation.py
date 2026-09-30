# -*- coding: utf-8 -*-
"""
Generate the IT Expert Testing & ISO/IEC 25010 Evaluation Form for A.E.G.I.S. Capstone.
Creates: docs/IT_Expert_Testing_and_ISO25010_Evaluation_Form.docx
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

# Brand Color Tokens
COLOR_PRIMARY_HEX = "1B4D2E"       # CLSU Forest Green
COLOR_SECONDARY_HEX = "B8860B"     # CLSU Gold / Ochre
COLOR_DARK_HEX = "1F2937"          # Dark Charcoal
COLOR_LIGHT_BG_HEX = "F3F4F6"      # Very Light Gray
COLOR_BORDER_HEX = "D1D5DB"        # Border Gray
COLOR_TEXT_MUTED = "4B5563"        # Muted Gray

COLOR_PRIMARY = RGBColor(27, 77, 46)
COLOR_SECONDARY = RGBColor(184, 134, 11)
COLOR_DARK = RGBColor(31, 41, 55)
COLOR_MUTED = RGBColor(75, 85, 99)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    """Set background color of a cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Set subtle, professional borders on a table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def make_callout(doc, text, bold_prefix="NOTE: ", border_color="1B4D2E", bg_color="F0FDF4"):
    """Create a stylized callout banner."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    set_cell_background(cell, bg_color)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
        <w:top w:val="none"/>
        <w:right w:val="none"/>
        <w:bottom w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run_bold = p.add_run(bold_prefix)
    run_bold.bold = True
    run_bold.font.name = 'Calibri'
    run_bold.font.size = Pt(9.5)
    run_bold.font.color.rgb = COLOR_PRIMARY
    
    run_text = p.add_run(text)
    run_text.font.name = 'Calibri'
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = COLOR_DARK

def add_header(doc):
    """Institutional Header."""
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.paragraph_format.space_before = Pt(0)
    r1 = p_inst.add_run("CENTRAL LUZON STATE UNIVERSITY\n")
    r1.font.name = 'Calibri'
    r1.font.size = Pt(11)
    r1.bold = True
    r1.font.color.rgb = COLOR_PRIMARY
    
    r2 = p_inst.add_run("OFFICE OF STUDENT AFFAIRS • SCHOLARSHIP & AID DIVISION\n")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(9.5)
    r2.bold = True
    r2.font.color.rgb = COLOR_SECONDARY
    
    r3 = p_inst.add_run("Science City of Muñoz, Nueva Ecija, Philippines\n")
    r3.font.name = 'Calibri'
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = COLOR_MUTED

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(12)
    p_line.paragraph_format.space_before = Pt(0)
    r_line = p_line.add_run("―" * 58)
    r_line.font.color.rgb = RGBColor(209, 213, 219)
    r_line.font.size = Pt(8)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("IT EXPERT SOFTWARE TESTING & EVALUATION FORM")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(15)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Technical Quality Assessment based on ISO/IEC 25010 Software Engineering Standards\n")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(10)
    r_sub.bold = True
    r_sub.font.color.rgb = COLOR_DARK

    r_sys = p_sub.add_run("Project: Academic Evaluation & Grant Integrity System (A.E.G.I.S.)")
    r_sys.font.name = 'Calibri'
    r_sys.font.size = Pt(9.5)
    r_sys.font.color.rgb = COLOR_MUTED

def add_evaluator_profile(doc):
    """Section for IT Expert Demographics & Background."""
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(12)
    p_sec.paragraph_format.space_after = Pt(6)
    r_sec = p_sec.add_run("I. IT EXPERT PROFILE & EVALUATOR CREDENTIALS")
    r_sec.font.name = 'Calibri'
    r_sec.font.size = Pt(11)
    r_sec.bold = True
    r_sec.font.color.rgb = COLOR_PRIMARY

    tbl = doc.add_table(rows=5, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)

    fields = [
        ("Evaluator Name (Optional / Confidential):", "Designation / Current Role:"),
        ("Highest Degree Attained:", "[  ] BS in IT / CS / CpE / IS   [  ] Master's Degree (MSIT/MIT/MCS)\n[  ] Doctorate Degree (Ph.D./DIT)   [  ] Other: ________________"),
        ("Area of Technical Specialization:", "[  ] Software Engineering / Architecture    [  ] Cybersecurity / InfoSec\n[  ] AI / Machine Learning / Forensics    [  ] Database / Cloud Systems\n[  ] Quality Assurance / Testing               [  ] IT Faculty / Academician"),
        ("Years of Professional IT Experience:", "[  ] 1 to 3 years    [  ] 4 to 6 years    [  ] 7 to 10 years    [  ] Over 10 years"),
        ("Professional Certifications / Affiliations:", "[  ] None    [  ] AWS / Cloud Certified    [  ] CISSP / CEH / CompTIA Security+\n[  ] Oracle / Database    [  ] Agile / Scrum Master    [  ] Other: ________________")
    ]

    col_widths = [Inches(2.5), Inches(4.0)]

    for row_idx, (label, default_val) in enumerate(fields):
        row = tbl.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = col_widths[0], col_widths[1]
        set_cell_margins(cell_lbl, top=70, bottom=70, left=100, right=100)
        set_cell_margins(cell_val, top=70, bottom=70, left=100, right=100)
        set_cell_background(cell_lbl, COLOR_LIGHT_BG_HEX)

        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(9)
        r0.bold = True
        r0.font.color.rgb = COLOR_DARK

        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(default_val)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_technical_test_matrix(doc):
    """Part II: Hands-On Technical Verification Matrix."""
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(12)
    p_sec.paragraph_format.space_after = Pt(4)
    r_sec = p_sec.add_run("II. HANDS-ON TECHNICAL VERIFICATION TEST MATRIX")
    r_sec.font.name = 'Calibri'
    r_sec.font.size = Pt(11)
    r_sec.bold = True
    r_sec.font.color.rgb = COLOR_PRIMARY

    make_callout(
        doc,
        "Instructions for Evaluator: Please conduct the following hands-on verification test scenarios across the A.E.G.I.S. system. Record your verdict for each scenario as [P] Pass, [F] Fail, or [NA] Not Applicable, along with observed execution latency and technical comments.",
        bold_prefix="EVALUATION PROTOCOL: ",
        border_color="1B4D2E",
        bg_color="F0FDF4"
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    test_cases = [
        ("TC-TECH-01", "Authentication & Cryptographic Security", 
         "Attempt SQL Injection & brute force bypass on `/login`. Verify Bcrypt hash (cost=12), rate-limiting middleware (5 attempts/min), and CSRF token binding.", 
         "SQL injection blocked; password hashes stored with irreversible Bcrypt; IP rate-limiter returns 429 Too Many Requests.",
         "[  ] Pass\n[  ] Fail"),
        
        ("TC-TECH-02", "Multi-Factor Authentication (MFA) & Hashing", 
         "Trigger OTP verification flow. Inspect database storage of `otp_code`. Verify SHA-256 zero-knowledge hashing at rest and 10-minute dynamic TTL expiry.", 
         "6-digit OTP stored as 64-char SHA-256 hash in DB; expired tokens rejected; brute-force locked out; universal demo code accepted for dummy accounts.",
         "[  ] Pass\n[  ] Fail"),

        ("TC-TECH-03", "Data Protection & Column-Level Encryption", 
         "Inspect database storage of sensitive student profile fields (e.g. bank account number, guardian details, and identity documents).", 
         "Sensitive columns stored using AES-256-CBC encryption; raw queries in DB return ciphertext; application dynamically decrypts for authorized sessions.",
         "[  ] Pass\n[  ] Fail"),

        ("TC-TECH-04", "Role-Based Access Control (RBAC) & Boundary Isolation", 
         "Authenticate as Student and attempt direct URL navigation to administrative routes (`/admin/dashboard`, `/superadmin/users`, `/superadmin/settings`).", 
         "HTTP 403 Forbidden or redirect to unauthorized notice; route middleware strictly isolates role boundaries; staff assignments restrict application scope.",
         "[  ] Pass\n[  ] Fail"),

        ("TC-TECH-05", "AI Multi-Detector Image Forensics Pipeline", 
         "Upload a Certificate of Grades (COG) with digitally altered grades (whiteout/clone-stamp). Inspect AI scan output, ELA heatmap, and OCR score.", 
         "Tesseract OCR extracts GWA; ELA detects compression inconsistencies; ORB clone detector flags copy-paste duplication; weighted fusion generates risk score.",
         "[  ] Pass\n[  ] Fail"),

        ("TC-TECH-06", "Tamper-Evident Audit Logging & Accountability", 
         "Execute administrative actions (approve application, modify system setting, export student data). Verify structured audit trail records.", 
         "Logs recorded in `admin_action_logs`, `config_change_logs`, and `export_access_logs` with actor ID, IP address, user agent, timestamp, and payload diff.",
         "[  ] Pass\n[  ] Fail"),

        ("TC-TECH-07", "Session Security & Cookie Hardening", 
         "Inspect HTTP response headers and cookie flags on authenticated HTTPS traffic (via DevTools Application/Network panel).", 
         "`Strict-Transport-Security`, `X-Frame-Options: SAMEORIGIN`, `X-Content-Type-Options: nosniff`, `CSP` headers present; session cookies flagged Secure & HttpOnly.",
         "[  ] Pass\n[  ] Fail"),

        ("TC-TECH-08", "Concurrency, Caching & Performance Efficiency", 
         "Simulate high concurrent page loads and document uploads. Test queue worker asynchronous processing for background AI scans.", 
         "Response compression (Gzip) active; N+1 queries eliminated via eager loading; heavy AI scan jobs handled asynchronously via queue worker without blocking UI.",
         "[  ] Pass\n[  ] Fail"),
    ]

    tbl = doc.add_table(rows=len(test_cases) + 1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)

    headers = ["ID", "Technical Test Scenario", "Execution Procedure & Inspection", "Expected Technical Standard", "Verdict / Obs."]
    col_widths = [Inches(0.9), Inches(1.5), Inches(2.0), Inches(1.4), Inches(0.7)]

    # Header Row
    hdr_row = tbl.rows[0]
    for c_idx, title in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = 'Calibri'
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, (tid, name, proc, exp, verd) in enumerate(test_cases):
        row = tbl.rows[r_idx + 1]
        bg_col = COLOR_LIGHT_BG_HEX if r_idx % 2 == 1 else "FFFFFF"

        vals = [tid, name, proc, exp, verd]
        for c_idx, text in enumerate(vals):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_background(cell, bg_col)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r.bold = True
            elif c_idx == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

def add_iso25010_section(doc):
    """Part III: ISO/IEC 25010 Software Quality Instrument."""
    doc.add_page_break()

    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(10)
    p_sec.paragraph_format.space_after = Pt(4)
    r_sec = p_sec.add_run("III. ISO/IEC 25010 SOFTWARE QUALITY EVALUATION INSTRUMENT")
    r_sec.font.name = 'Calibri'
    r_sec.font.size = Pt(11)
    r_sec.bold = True
    r_sec.font.color.rgb = COLOR_PRIMARY

    # Rating Scale Table
    scale_tbl = doc.add_table(rows=6, cols=3)
    scale_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    scale_tbl.autofit = False
    set_table_borders(scale_tbl)

    scale_headers = ["Rating Scale", "Statistical Range", "Qualitative Description & Standard"]
    scale_data = [
        ("5 - Strongly Agree (SA)", "4.20 - 5.00", "Exemplary implementation; exceeds industry benchmarks with zero architectural defects."),
        ("4 - Agree (A)", "3.40 - 4.19", "Robust implementation; meets all standard technical specifications with negligible observations."),
        ("3 - Moderately Agree (MA)", "2.60 - 3.39", "Acceptable implementation; satisfies fundamental requirements but has minor technical room for optimization."),
        ("2 - Disagree (D)", "1.80 - 2.59", "Substandard implementation; features identifiable architectural shortcomings or vulnerability risks."),
        ("1 - Strongly Disagree (SD)", "1.00 - 1.79", "Critically deficient implementation; fails to meet baseline software quality and security standards.")
    ]

    s_widths = [Inches(1.8), Inches(1.3), Inches(3.4)]

    # Header
    for c_idx, title in enumerate(scale_headers):
        cell = scale_tbl.rows[0].cells[c_idx]
        cell.width = s_widths[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        set_cell_background(cell, COLOR_SECONDARY_HEX)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = 'Calibri'
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, (scale, rng, desc) in enumerate(scale_data):
        row = scale_tbl.rows[r_idx + 1]
        for c_idx, txt in enumerate([scale, rng, desc]):
            cell = row.cells[c_idx]
            cell.width = s_widths[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else COLOR_LIGHT_BG_HEX)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            r.font.name = 'Calibri'
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx < 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 8 ISO Categories
    categories = [
        ("1. Functional Suitability", [
            ("FS-01", "Functional Completeness", "The system covers all specified scholarship management tasks, including application submission, multi-tiered document evaluation, status tracking, and announcement distribution."),
            ("FS-02", "Functional Correctness", "The system executes all computation routines accurately, including GWA pre-screening, income thresholds, academic standing validation, and AI risk scoring."),
            ("FS-03", "Functional Appropriateness", "The implemented technical features directly facilitate and streamline administrative scholarship workflows without extraneous or redundant operations.")
        ]),
        ("2. Performance Efficiency", [
            ("PE-01", "Time Behaviour", "System response times for standard transactions (page renders, application queries, state updates) consistently execute within acceptable latency thresholds (< 1.5 seconds)."),
            ("PE-02", "Resource Utilization", "Server CPU, RAM, and database I/O resources are consumed efficiently through eager loading, query optimization, response compression (Gzip), and OPcache execution."),
            ("PE-03", "Capacity & Throughput", "The system effectively manages concurrent applicant uploads and background worker queuing without transaction deadlocks or performance degradation.")
        ]),
        ("3. Compatibility", [
            ("CO-01", "Co-existence", "The software operates stably in shared hosting environments and multi-container architectures without conflicting with co-located services or system daemons."),
            ("CO-02", "Interoperability", "The application integrates seamlessly with external web services and APIs (e.g. Brevo SMTP email delivery, Cloudflare R2 object storage, HuggingFace AI endpoints).")
        ]),
        ("4. Usability", [
            ("US-01", "Appropriateness Recognisability", "The user interface utilizes intuitive design patterns, breadcrumbs, status badges, and semantic layouts that allow users to readily understand system capabilities."),
            ("US-02", "Learnability", "New students, OSA staff, and administrative evaluators can operate the application with minimal onboarding training via intuitive form wizards and contextual tips."),
            ("US-03", "User Error Protection", "Interactive inputs feature comprehensive client-side and server-side validation rules, unified reconfirmation dialogues (SweetAlert2), and clear error feedback to mitigate user mistakes.")
        ]),
        ("5. Reliability", [
            ("RE-01", "Maturity", "The system demonstrates high operational stability, passing comprehensive automated unit/feature test suites (250+ test cases) without unhandled exceptions."),
            ("RE-02", "Fault Tolerance", "The system gracefully handles external dependency interruptions (e.g. AI cold start backoff, email failover drivers, and database busy retries) without crashing."),
            ("RE-03", "Recoverability", "In the event of network disruption or transaction failure, the system preserves state integrity and allows rapid resumption of interrupted workflows.")
        ]),
        ("6. Security", [
            ("SE-01", "Confidentiality", "Unauthorized access to private student data is strictly prevented through role-based access control, cryptographic session cookies, and AES-256 column-level database encryption."),
            ("SE-02", "Integrity", "The system prevents unauthorized modification of applicant records, evaluation scores, and audit trails through CSRF protection, signed URLs, and tamper-evident logging."),
            ("SE-03", "Non-repudiation", "System transactions (approvals, rejections, settings updates, user deletions) are immutably tied to the executing identity with IP, user agent, and timestamp metadata."),
            ("SE-04", "Authenticity", "User identity is robustly confirmed through cryptographically secure Multi-Factor Authentication (MFA OTP) with SHA-256 hash storage at rest and brute-force mitigation.")
        ]),
        ("7. Maintainability", [
            ("MA-01", "Modularity", "The codebase is structured according to clear architectural boundaries (Controllers, Services, Jobs, Models, Middleware) ensuring changes to one module have minimal unintended impacts."),
            ("MA-02", "Reusability", "Core components (alert systems, forensic analyzers, notification broadcasters, PDF generators) are abstracted into modular services and Blade components for reuse."),
            ("MA-03", "Analysability & Testability", "The application provides clear diagnostic logging, audit trails, and automated test fixtures enabling rapid troubleshooting and regression testing.")
        ]),
        ("8. Portability", [
            ("PO-01", "Adaptability", "The application adapts seamlessly across modern web browsers (Chrome, Edge, Firefox, Safari) and screen sizes (Desktop, Tablet, Mobile) with fluid responsiveness."),
            ("PO-02", "Installability", "The application deployment pipeline is standardized via containerization (Docker), environment configuration (`.env`), and automated database migration/seeding scripts.")
        ]),
    ]

    for cat_title, items in categories:
        p_cat = doc.add_paragraph()
        p_cat.paragraph_format.space_before = Pt(8)
        p_cat.paragraph_format.space_after = Pt(2)
        r_cat = p_cat.add_run(cat_title)
        r_cat.font.name = 'Calibri'
        r_cat.font.size = Pt(10)
        r_cat.bold = True
        r_cat.font.color.rgb = COLOR_PRIMARY

        tbl = doc.add_table(rows=len(items) + 1, cols=7)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders(tbl)

        t_widths = [Inches(0.8), Inches(1.5), Inches(2.7), Inches(0.3), Inches(0.3), Inches(0.3), Inches(0.3), Inches(0.3)]

        # Header
        hdr = tbl.rows[0]
        hdr_labels = ["Code", "Sub-Characteristic", "Technical Quality Evaluation Statement", "5", "4", "3", "2", "1"]
        for c_idx, lbl in enumerate(hdr_labels[:7]):
            cell = hdr.cells[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            set_cell_background(cell, COLOR_PRIMARY_HEX)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(lbl)
            r.font.name = 'Calibri'
            r.font.size = Pt(8)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            if c_idx >= 3 or c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r_idx, (code, subc, stmt) in enumerate(items):
            row = tbl.rows[r_idx + 1]
            bg_col = "FFFFFF" if r_idx % 2 == 0 else COLOR_LIGHT_BG_HEX
            for c_idx in range(7):
                cell = row.cells[c_idx]
                set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
                set_cell_background(cell, bg_col)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = p.add_run(code)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(7.5)
                    r.bold = True
                    r.font.color.rgb = COLOR_DARK
                elif c_idx == 1:
                    r = p.add_run(subc)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(7.5)
                    r.bold = True
                    r.font.color.rgb = COLOR_DARK
                elif c_idx == 2:
                    r = p.add_run(stmt)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(7.5)
                    r.font.color.rgb = COLOR_DARK
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = p.add_run("[   ]")
                    r.font.name = 'Calibri'
                    r.font.size = Pt(7.5)
                    r.font.color.rgb = COLOR_MUTED

def add_qualitative_assessment(doc):
    """Part IV: Qualitative Assessment, Recommendations & Endorsement."""
    doc.add_page_break()

    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(10)
    p_sec.paragraph_format.space_after = Pt(6)
    r_sec = p_sec.add_run("IV. QUALITATIVE TECHNICAL ASSESSMENT & REMARKS")
    r_sec.font.name = 'Calibri'
    r_sec.font.size = Pt(11)
    r_sec.bold = True
    r_sec.font.color.rgb = COLOR_PRIMARY

    prompts = [
        ("1. Key Architectural & Technical Strengths of the System:",
         "What aspects of the system's software architecture, security mechanisms (MFA/AES), or AI forensic pipeline stand out as commendable?"),
        
        ("2. Areas for Technical Refinement, Scalability, or Security Hardening:",
         "What optimizations, architectural improvements, or additional security controls do you suggest prior to widespread university deployment?"),

        ("3. Observations on AI Tamper Detection Accuracy & Usability in Academic Administration:",
         "How viable is the Explainable Forensic Decision Framework (EFDF) and multi-detector pipeline for aiding human evaluators at OSA?")
    ]

    for title, desc in prompts:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(8)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(title)
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(9.5)
        r_t.bold = True
        r_t.font.color.rgb = COLOR_DARK

        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(4)
        r_d = p_d.add_run(desc)
        r_d.font.name = 'Calibri'
        r_d.font.size = Pt(8.5)
        r_d.italic = True
        r_d.font.color.rgb = COLOR_MUTED

        # Box for comments
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        set_cell_background(cell, COLOR_LIGHT_BG_HEX)
        set_table_borders(tbl, color=COLOR_BORDER_HEX)

        p_box = cell.paragraphs[0]
        p_box.paragraph_format.space_before = Pt(0)
        p_box.paragraph_format.space_after = Pt(0)
        r_box = p_box.add_run("\n\n\n\n")
        r_box.font.size = Pt(9)

    # Deployment Recommendation & Sign-off
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_before = Pt(14)
    p_rec.paragraph_format.space_after = Pt(4)
    r_rec = p_rec.add_run("V. OVERALL TECHNICAL VERDICT & DEPLOYMENT ENDORSEMENT")
    r_rec.font.name = 'Calibri'
    r_rec.font.size = Pt(11)
    r_rec.bold = True
    r_rec.font.color.rgb = COLOR_PRIMARY

    p_chk = doc.add_paragraph()
    p_chk.paragraph_format.space_before = Pt(2)
    p_chk.paragraph_format.space_after = Pt(12)
    r_chk = p_chk.add_run(
        "[   ] FULLY ENDORSED: The system satisfies all ISO/IEC 25010 software quality benchmarks and is ready for institutional deployment.\n"
        "[   ] CONDITIONALLY ENDORSED: The system meets fundamental requirements; minor technical refinements recommended before full deployment.\n"
        "[   ] NOT RECOMMENDED: Significant architectural or security vulnerabilities exist that must be remediated prior to reconsidering deployment."
    )
    r_chk.font.name = 'Calibri'
    r_chk.font.size = Pt(9)
    r_chk.font.color.rgb = COLOR_DARK

    # Signature Block
    sig_tbl = doc.add_table(rows=2, cols=2)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False
    set_table_borders(sig_tbl, val="none")

    col_widths = [Inches(3.2), Inches(3.2)]

    cell_sig0 = sig_tbl.rows[0].cells[0]
    cell_sig1 = sig_tbl.rows[0].cells[1]
    cell_sig0.width, cell_sig1.width = col_widths[0], col_widths[1]

    p_s0 = cell_sig0.paragraphs[0]
    p_s0.paragraph_format.space_before = Pt(20)
    p_s0.paragraph_format.space_after = Pt(0)
    r_s0 = p_s0.add_run("_________________________________________\nEvaluator Signature Over Printed Name")
    r_s0.font.name = 'Calibri'
    r_s0.font.size = Pt(9)
    r_s0.bold = True

    p_s1 = cell_sig1.paragraphs[0]
    p_s1.paragraph_format.space_before = Pt(20)
    p_s1.paragraph_format.space_after = Pt(0)
    r_s1 = p_s1.add_run("_________________________________________\nDate of Technical Evaluation")
    r_s1.font.name = 'Calibri'
    r_s1.font.size = Pt(9)
    r_s1.bold = True

def main():
    doc = docx.Document()
    
    # Page Setup: Standard Letter, 0.75-inch margins
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    add_header(doc)
    add_evaluator_profile(doc)
    add_technical_test_matrix(doc)
    add_iso25010_section(doc)
    add_qualitative_assessment(doc)

    output_path = os.path.join("docs", "IT_Expert_Testing_and_ISO25010_Evaluation_Form.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path} ({os.path.getsize(output_path):,} bytes)")

if __name__ == "__main__":
    main()
