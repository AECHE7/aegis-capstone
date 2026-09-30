# -*- coding: utf-8 -*-
"""
Generate Role-Specific UAT Test Documents & ISO/IEC 25010:2023 Questionnaires
strictly based on docs/Client_Testing_and_ISO25010_End_User_Evaluation.docx
for A.E.G.I.S. Capstone 2 (Central Luzon State University - Office of Student Affairs).

Roles generated:
1. docs/UAT_Test_Script_Student_Role.docx
2. docs/UAT_Test_Script_Staff_Role.docx
3. docs/UAT_Test_Script_Admin_Role.docx
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

# Reference document colors and styling
SHD_BLUE = "D9EAF7"      # Light Soft Blue header for Test Scenarios and ISO items
SHD_GREEN = "E2F0D9"     # Light Soft Green header for Revision Log and Scoring
SHD_LIGHT_ROW = "F9FAFB" # Very light alternating row
BORDER_COLOR = "B0C4DE"  # Soft slate border

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Set inner padding for table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, fill_hex):
    """Set background color of a cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    """Set clean borders on table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def build_part1_header(doc, role_title, role_subtitle):
    """Build Part 1 Institutional letterhead and Title matching the reference doc."""
    p_clsu = doc.add_paragraph()
    p_clsu.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_clsu.paragraph_format.space_before = Pt(0)
    p_clsu.paragraph_format.space_after = Pt(2)
    
    r_uni = p_clsu.add_run("CENTRAL LUZON STATE UNIVERSITY\n")
    r_uni.bold = True
    r_uni.font.name = "Arial"
    r_uni.font.size = Pt(11)
    
    r_dept = p_clsu.add_run("College of Engineering • Department of Information Technology\n")
    r_dept.font.name = "Arial"
    r_dept.font.size = Pt(9.5)
    
    r_addr = p_clsu.add_run("Science City of Muñoz, Nueva Ecija, Philippines\n\n")
    r_addr.font.name = "Arial"
    r_addr.font.size = Pt(9)
    
    r_title = p_clsu.add_run("CLIENT SYSTEM TESTING AND ACCEPTANCE FORM\n")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(12)
    
    r_role = p_clsu.add_run(f"[{role_title.upper()}]\n")
    r_role.bold = True
    r_role.font.name = "Arial"
    r_role.font.size = Pt(10.5)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run(f"{role_subtitle} — Capstone Software Validation")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True

    p_purp = doc.add_paragraph()
    p_purp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_purp.paragraph_format.space_before = Pt(0)
    p_purp.paragraph_format.space_after = Pt(4)
    r_purp_lbl = p_purp.add_run("Purpose. ")
    r_purp_lbl.bold = True
    r_purp_lbl.font.name = "Arial"
    r_purp_lbl.font.size = Pt(9)
    r_purp_txt = p_purp.add_run(
        "This form documents the client's actual testing of the developed system. The tester should perform representative "
        "tasks and verify whether the major functions and outputs work as expected. Issues identified during testing should be "
        "corrected and, when necessary, subjected to retesting."
    )
    r_purp_txt.font.name = "Arial"
    r_purp_txt.font.size = Pt(9)

    # Section A: Instructions
    p_inst_hdr = doc.add_paragraph()
    p_inst_hdr.paragraph_format.space_before = Pt(5)
    p_inst_hdr.paragraph_format.space_after = Pt(3)
    p_inst_hdr.paragraph_format.keep_with_next = True
    r_ih = p_inst_hdr.add_run("A. TESTING INSTRUCTIONS")
    r_ih.bold = True
    r_ih.font.name = "Arial"
    r_ih.font.size = Pt(10)

    instructions = [
        "The development team shall briefly orient the tester, then allow the representative to perform the actual test.",
        "Test the major modules, workflows, and outputs relevant to the role's intended use.",
        "Mark each item as PASS, FAIL, NEEDS REVISION, or N/A.",
        "Record errors, missing functions, confusing steps, inaccurate outputs, or other concerns in the Remarks column.",
        "Items marked FAIL or NEEDS REVISION should be corrected and retested before final acceptance."
    ]
    for inst in instructions:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(inst)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)

def build_part1_metadata_table(doc, role_default, tester_name=""):
    """Metadata block matching Table 0 of reference document."""
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    meta_items = [
        ("Project / System Title: A.E.G.I.S. (Automated Evaluation & Grade Integrity System)",
         "Client / Organization: Central Luzon State University — Office of Student Affairs (CLSU OSA)"),
        (f"Tester / Client Representative: {tester_name}",
         f"Role / Designation: {role_default}"),
        ("Date of Testing: October 2026",
         "Environment / Build: Cloud Staging Portal (https://clsu.osa.scholarship) v1.0.0")
    ]

    for r_i, (c0_text, c1_text) in enumerate(meta_items):
        row = table.rows[r_i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(3.7)
        c1.width = Inches(3.7)
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(c0_text)
        r0.font.name = "Arial"
        r0.font.size = Pt(8.5)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(c1_text)
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_part1_scenarios_table(doc, scenarios):
    """Build Table B matching the 7-column reference layout."""
    p_sc_hdr = doc.add_paragraph()
    p_sc_hdr.paragraph_format.space_before = Pt(6)
    p_sc_hdr.paragraph_format.space_after = Pt(3)
    p_sc_hdr.paragraph_format.keep_with_next = True
    r_sh = p_sc_hdr.add_run("B. CLIENT TEST SCENARIOS")
    r_sh.bold = True
    r_sh.font.name = "Arial"
    r_sh.font.size = Pt(10)

    table = doc.add_table(rows=len(scenarios) + 1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    headers = ["No.", "Module / Feature", "Task / Test Scenario", "Expected Result", "Actual Result", "Status", "Remarks"]
    widths = [Inches(0.35), Inches(1.1), Inches(1.5), Inches(1.4), Inches(1.15), Inches(1.0), Inches(0.9)]

    # Header row
    for j, h in enumerate(headers):
        c = table.cell(0, j)
        c.width = widths[j]
        set_cell_shading(c, SHD_BLUE)
        set_cell_margins(c, 70, 70, 60, 60)
        p = c.paragraphs[0]
        if j in [0, 5]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)

    for i, sc in enumerate(scenarios, 1):
        row = table.rows[i]
        cols_data = [
            str(i),
            sc['module'],
            sc['task'],
            sc['expected'],
            sc.get('actual', ''),
            "☐ Pass\n☐ Fail\n☐ Needs Rev.\n☐ N/A",
            sc.get('remarks', '')
        ]

        for j, val in enumerate(cols_data):
            c = row.cells[j]
            c.width = widths[j]
            set_cell_margins(c, 50, 50, 60, 60)
            if i % 2 == 1:
                set_cell_shading(c, SHD_LIGHT_ROW)

            p = c.paragraphs[0]
            if j == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            if j == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_part1_issue_log_and_result(doc, role_label="Client Representative"):
    """Build Section C (Issue Log), Section D (Testing Result), and Signatures."""
    p_iss_hdr = doc.add_paragraph()
    p_iss_hdr.paragraph_format.space_before = Pt(6)
    p_iss_hdr.paragraph_format.space_after = Pt(3)
    p_iss_hdr.paragraph_format.keep_with_next = True
    r_ih = p_iss_hdr.add_run("C. ISSUE / REVISION LOG")
    r_ih.bold = True
    r_ih.font.name = "Arial"
    r_ih.font.size = Pt(10)

    # Issue Table: 5 rows
    table_iss = doc.add_table(rows=5, cols=6)
    table_iss.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_iss)

    iss_headers = ["No.", "Issue / Observation", "Required Revision / Action", "Priority", "Retest Result", "Remarks"]
    iss_widths = [Inches(0.35), Inches(2.0), Inches(2.0), Inches(1.1), Inches(1.0), Inches(0.95)]

    for j, h in enumerate(iss_headers):
        c = table_iss.cell(0, j)
        c.width = iss_widths[j]
        set_cell_shading(c, SHD_GREEN)
        set_cell_margins(c, 60, 60, 60, 60)
        p = c.paragraphs[0]
        if j in [0, 3, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)

    for r_i in range(1, 5):
        row = table_iss.rows[r_i]
        for c_i in range(6):
            c = row.cells[c_i]
            c.width = iss_widths[c_i]
            set_cell_margins(c, 50, 50, 60, 60)
            p = c.paragraphs[0]
            if c_i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(str(r_i))
                r.font.name = "Arial"
                r.font.size = Pt(8)
                r.bold = True
            elif c_i == 3:
                r = p.add_run("☐ High\n☐ Med\n☐ Low")
                r.font.name = "Arial"
                r.font.size = Pt(7.5)
            elif c_i == 4:
                r = p.add_run("☐ Passed\n☐ For Retest")
                r.font.name = "Arial"
                r.font.size = Pt(7.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Section D: Result
    p_res_hdr = doc.add_paragraph()
    p_res_hdr.paragraph_format.space_before = Pt(6)
    p_res_hdr.paragraph_format.space_after = Pt(2)
    p_res_hdr.paragraph_format.keep_with_next = True
    r_rh = p_res_hdr.add_run("D. CLIENT TESTING RESULT")
    r_rh.bold = True
    r_rh.font.name = "Arial"
    r_rh.font.size = Pt(10)

    p_base = doc.add_paragraph()
    p_base.paragraph_format.space_before = Pt(0)
    p_base.paragraph_format.space_after = Pt(2)
    r_b = p_base.add_run("Based on the testing performed, the system is:")
    r_b.font.name = "Arial"
    r_b.font.size = Pt(8.5)

    results = [
        "☐ ACCEPTED — major required functions operated satisfactorily and no critical issue prevents intended use.",
        "☐ ACCEPTED WITH MINOR REVISIONS — the system is generally usable, subject to the corrections listed above.",
        "☐ FOR REVISION AND RETESTING — one or more significant issues must be corrected before acceptance."
    ]
    for res in results:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(res)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)

    p_comm = doc.add_paragraph()
    p_comm.paragraph_format.space_before = Pt(4)
    p_comm.paragraph_format.space_after = Pt(10)
    p_comm.add_run("Client Comments / Recommendations:\n" + "_" * 95 + "\n" + "_" * 95).font.size = Pt(8.5)

    # Signatures
    table_sign = doc.add_table(rows=1, cols=2)
    table_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_sign, val="none")

    c0, c1 = table_sign.cell(0, 0), table_sign.cell(0, 1)
    c0.width = Inches(3.7)
    c1.width = Inches(3.7)

    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(6)
    p0.add_run("_________________________________________\n").bold = True
    p0.add_run(f"{role_label} Signature over Printed Name\nDate: ________________________").font.size = Pt(8.5)

    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(6)
    p1.add_run("_________________________________________\n").bold = True
    p1.add_run("JOSHUA RAZON / NORIEL GADIANO / JOHN ANDREI CARILLO II\nStudent Researchers / Project Leaders\nDate: ________________________").font.size = Pt(8.5)

def build_part2_iso_questionnaire(doc, role_title, respondent_role_box, statements_dict):
    """Build Part 2 ISO/IEC 25010:2023 Evaluation Form matching the reference document."""
    doc.add_page_break()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_t = p_title.add_run("END-USER SYSTEM EVALUATION FORM\n")
    r_t.bold = True
    r_t.font.name = "Arial"
    r_t.font.size = Pt(11.5)

    r_st = p_title.add_run("ISO/IEC 25010:2023-Aligned Product Quality Evaluation\n")
    r_st.font.name = "Arial"
    r_st.font.size = Pt(9.5)

    r_rt = p_title.add_run(f"Evaluation Instrument — {role_title}")
    r_rt.bold = True
    r_rt.font.name = "Arial"
    r_rt.font.size = Pt(9)
    r_rt.font.italic = True

    p_purp = doc.add_paragraph()
    p_purp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_purp.paragraph_format.space_before = Pt(4)
    p_purp.paragraph_format.space_after = Pt(3)
    p_purp.add_run("Purpose. ").bold = True
    p_purp.add_run(
        "This questionnaire gathers end-user feedback on system quality after actual use or guided testing. The statements "
        "are written in end-user language and aligned with relevant ISO/IEC 25010:2023 product quality characteristics. "
        "Use N/A when a statement is not applicable or cannot reasonably be evaluated."
    ).font.size = Pt(8.5)

    p_priv = doc.add_paragraph()
    p_priv.paragraph_format.space_before = Pt(3)
    p_priv.paragraph_format.space_after = Pt(2)
    p_priv.paragraph_format.keep_with_next = True
    p_priv.add_run("PRIVACY AND VOLUNTARY PARTICIPATION NOTICE").bold = True
    p_priv.runs[0].font.size = Pt(9)

    p_priv_txt = doc.add_paragraph()
    p_priv_txt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_priv_txt.paragraph_format.space_before = Pt(0)
    p_priv_txt.paragraph_format.space_after = Pt(4)
    p_priv_txt.add_run(
        "Participation is voluntary. Responses will be used only for system evaluation, academic documentation, and project "
        "improvement. Personal information, if collected, should be limited to what is necessary and should not be disclosed to "
        "unauthorized persons in accordance with R.A. 10173 (Data Privacy Act of 2012). Optional profile fields may be left blank."
    ).font.size = Pt(8.5)

    p_scale = doc.add_paragraph()
    p_scale.paragraph_format.space_before = Pt(3)
    p_scale.paragraph_format.space_after = Pt(2)
    p_scale.paragraph_format.keep_with_next = True
    p_scale.add_run("RATING SCALE").bold = True
    p_scale.runs[0].font.size = Pt(9)

    p_scale_txt = doc.add_paragraph()
    p_scale_txt.paragraph_format.space_before = Pt(0)
    p_scale_txt.paragraph_format.space_after = Pt(2)
    p_scale_txt.add_run(
        "5 - Strongly Agree     4 - Agree     3 - Neither Agree nor Disagree     2 - Disagree     1 - Strongly Disagree     N/A - Not Applicable / Cannot Evaluate"
    ).font.size = Pt(8)

    p_dir = doc.add_paragraph()
    p_dir.paragraph_format.space_before = Pt(0)
    p_dir.paragraph_format.space_after = Pt(4)
    r_dir = p_dir.add_run("Direction: After using the system, check one rating for each statement.")
    r_dir.bold = True
    r_dir.font.size = Pt(8.5)

    # Respondent info table
    tbl_resp = doc.add_table(rows=2, cols=4)
    tbl_resp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_resp)

    r_items = [
        ("Project / System Title", "A.E.G.I.S. (Automated Evaluation & Grade Integrity System)", "Date of Evaluation", "October 2026"),
        ("Respondent Role / User Type", respondent_role_box, "Organization / Office", "Central Luzon State University — Office of Student Affairs (CLSU OSA)")
    ]
    r_widths = [Inches(1.8), Inches(2.2), Inches(1.5), Inches(1.9)]
    for r_i, (k0, v0, k1, v1) in enumerate(r_items):
        row = tbl_resp.rows[r_i]
        for c_i, text in enumerate([k0, v0, k1, v1]):
            c = row.cells[c_i]
            c.width = r_widths[c_i]
            set_cell_margins(c, 50, 50, 60, 60)
            if c_i in [0, 2]:
                set_cell_shading(c, "F3F4F6")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            if c_i in [0, 2]:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 26 ISO Statements Table
    total_rows = 1 + len(statements_dict) + sum(len(items) for items in statements_dict.values())
    tbl_iso = doc.add_table(rows=total_rows, cols=3)
    tbl_iso.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_iso)

    col_widths = [Inches(0.45), Inches(4.7), Inches(2.25)]

    # Header row
    for j, h in enumerate(["No.", "Evaluation Statement", "Rating"]):
        c = tbl_iso.cell(0, j)
        c.width = col_widths[j]
        set_cell_shading(c, SHD_BLUE)
        set_cell_margins(c, 60, 60, 60, 60)
        p = c.paragraphs[0]
        if j in [0, 2]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)

    curr_r = 1
    q_num = 1
    for dim_title, items in statements_dict.items():
        # Dimension category row
        c_dim = tbl_iso.cell(curr_r, 0)
        c_dim.merge(tbl_iso.cell(curr_r, 1))
        c_dim.merge(tbl_iso.cell(curr_r, 2))
        set_cell_shading(c_dim, SHD_BLUE)
        set_cell_margins(c_dim, 50, 50, 80, 80)
        p_d = c_dim.paragraphs[0]
        p_d.paragraph_format.space_after = Pt(0)
        r_d = p_d.add_run(dim_title.upper())
        r_d.bold = True
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8.5)
        curr_r += 1

        for stmt in items:
            row = tbl_iso.rows[curr_r]
            
            c0 = row.cells[0]
            c0.width = col_widths[0]
            set_cell_margins(c0, 40, 40, 50, 50)
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(str(q_num))
            r0.font.name = "Arial"
            r0.font.size = Pt(8)
            r0.bold = True

            c1 = row.cells[1]
            c1.width = col_widths[1]
            set_cell_margins(c1, 40, 40, 60, 60)
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(stmt)
            r1.font.name = "Arial"
            r1.font.size = Pt(8)

            c2 = row.cells[2]
            c2.width = col_widths[2]
            set_cell_margins(c2, 40, 40, 50, 50)
            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run("☐ 5   ☐ 4   ☐ 3   ☐ 2   ☐ 1   ☐ N/A")
            r2.font.name = "Arial"
            r2.font.size = Pt(7.5)

            curr_r += 1
            q_num += 1

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Section B: Overall Assessment
    p_ov_hdr = doc.add_paragraph()
    p_ov_hdr.paragraph_format.space_before = Pt(6)
    p_ov_hdr.paragraph_format.space_after = Pt(2)
    p_ov_hdr.paragraph_format.keep_with_next = True
    r_oh = p_ov_hdr.add_run("OVERALL ASSESSMENT")
    r_oh.bold = True
    r_oh.font.name = "Arial"
    r_oh.font.size = Pt(9.5)

    tbl_ov = doc.add_table(rows=3, cols=2)
    tbl_ov.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_ov)

    ov_stmts = [
        "Overall, I am satisfied with the system.",
        "I would be comfortable using this system for its intended purpose.",
        "The system is ready for actual use, subject to necessary revisions."
    ]
    for idx, ost in enumerate(ov_stmts):
        row = tbl_ov.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(5.15)
        c1.width = Inches(2.25)
        set_cell_margins(c0, 40, 40, 60, 60)
        set_cell_margins(c1, 40, 40, 50, 50)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(ost)
        r0.font.name = "Arial"
        r0.font.size = Pt(8)

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run("☐ 5   ☐ 4   ☐ 3   ☐ 2   ☐ 1")
        r1.font.name = "Arial"
        r1.font.size = Pt(7.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Comments & Recommendations
    p_cr_hdr = doc.add_paragraph()
    p_cr_hdr.paragraph_format.space_before = Pt(6)
    p_cr_hdr.paragraph_format.space_after = Pt(2)
    p_cr_hdr.paragraph_format.keep_with_next = True
    r_cr = p_cr_hdr.add_run("COMMENTS AND RECOMMENDATIONS (Respondent)")
    r_cr.bold = True
    r_cr.font.name = "Arial"
    r_cr.font.size = Pt(9.5)

    comms = [
        "Features or aspects I liked:  " + "_" * 85,
        "Problems or difficulties I encountered:  " + "_" * 78,
        "Suggested improvements:  " + "_" * 84
    ]
    for cm in comms:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(cm)
        r.font.name = "Arial"
        r.font.size = Pt(8)

    # Statistical Scoring Guide Table
    p_sc_hdr = doc.add_paragraph()
    p_sc_hdr.paragraph_format.space_before = Pt(6)
    p_sc_hdr.paragraph_format.space_after = Pt(2)
    p_sc_hdr.paragraph_format.keep_with_next = True
    r_sch = p_sc_hdr.add_run("RESEARCHER / INSTRUCTOR SCORING GUIDE")
    r_sch.bold = True
    r_sch.font.name = "Arial"
    r_sch.font.size = Pt(9.5)

    tbl_sc = doc.add_table(rows=6, cols=2)
    tbl_sc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sc)

    ranges = [
        ("Mean Range", "Interpretation"),
        ("4.21–5.00", "Strongly Agree / Very High Acceptability"),
        ("3.41–4.20", "Agree / High Acceptability"),
        ("2.61–3.40", "Neither Agree nor Disagree / Moderate Acceptability"),
        ("1.81–2.60", "Disagree / Low Acceptability"),
        ("1.00–1.80", "Strongly Disagree / Very Low Acceptability")
    ]
    for r_i, (rng, interp) in enumerate(ranges):
        row = tbl_sc.rows[r_i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(5.2)
        set_cell_margins(c0, 40, 40, 60, 60)
        set_cell_margins(c1, 40, 40, 60, 60)
        if r_i == 0:
            set_cell_shading(c0, "D9EAD3")
            set_cell_shading(c1, "D9EAD3")

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(rng)
        r0.font.name = "Arial"
        r0.font.size = Pt(8)
        if r_i == 0:
            r0.bold = True

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(interp)
        r1.font.name = "Arial"
        r1.font.size = Pt(8)
        if r_i == 0:
            r1.bold = True

    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_before = Pt(4)
    p_ref.paragraph_format.space_after = Pt(0)
    r_rf = p_ref.add_run("Reference note: ISO/IEC 25010:2023, Systems and software engineering — Systems and software Quality Requirements and Evaluation (SQuaRE) — Product quality model.")
    r_rf.font.name = "Arial"
    r_rf.font.size = Pt(7.5)
    r_rf.font.italic = True


# ==============================================================================
# 1. BUILD STUDENT ROLE DOC
# ==============================================================================
def create_student_document():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.55)
        s.bottom_margin = Inches(0.55)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    build_part1_header(
        doc,
        role_title="STUDENT APPLICANT ROLE",
        role_subtitle="Undergraduate Student End-User Validation"
    )

    build_part1_metadata_table(
        doc,
        role_default="Student Applicant (Undergraduate / @clsu2.edu.ph)"
    )

    student_scenarios = [
        {
            "module": "Student Registration & Security",
            "task": "Student registers account using institutional '@clsu2.edu.ph' webmail and sets strong password.",
            "expected": "System validates institutional domain, blocks non-CLSU emails, hashes credentials, and triggers verification notice.",
        },
        {
            "module": "Profile & GWA Setup",
            "task": "Student fills in demographic info, college/course, year level, contact info, and academic GWA.",
            "expected": "GWA is validated (1.00 to 5.00), profile completeness meter updates to 100%, and data persists accurately.",
        },
        {
            "module": "Catalog Discovery & Filtering",
            "task": "Student explores available scholarships and filters by criteria (minimum GWA, eligible colleges, category).",
            "expected": "Catalog updates dynamically; eligible scholarships show active 'Apply Now' buttons; ineligible programs show reasons.",
        },
        {
            "module": "Application & Custom Fields",
            "task": "Student applies for scholarship and completes custom questionnaire fields (income tier, occupation, essays).",
            "expected": "Dynamic fields render correctly (text, dropdown, file); client draft auto-saves; required validator prevents empty submits.",
        },
        {
            "module": "Document Upload & Preview",
            "task": "Student uploads required documents (Certificate of Grades and Certificate of Registration) in PDF/PNG/JPEG.",
            "expected": "File format and size (<10MB) validated; thumbnail preview opens for pre-submission verification.",
        },
        {
            "module": "Submission & Progress Tracker",
            "task": "Student reviews submission summary, agrees to R.A. 10173 data privacy terms, and submits application.",
            "expected": "Confirmation alert generates unique tracking ID; application status moves to 'Pending Review' on dashboard timeline.",
        },
        {
            "module": "Deficiency Resubmission",
            "task": "Student opens application marked 'Returned for Correction', views staff remarks, and resubmits corrected file.",
            "expected": "Staff deficiency instructions display in amber alert; flagged field unlocks for re-upload; status updates to 'Resubmitted'.",
        },
        {
            "module": "Notifications & Official PDF",
            "task": "Student checks status change notifications and downloads official approved application form with QR seal.",
            "expected": "In-app and email notifications arrive promptly; certified PDF downloads with official CLSU OSA seal and QR verification.",
        }
    ]
    build_part1_scenarios_table(doc, student_scenarios)
    build_part1_issue_log_and_result(doc, role_label="Student Applicant Representative")

    # 26 ISO Statements contextualized for student
    student_iso_stmts = {
        "1. Functional Suitability": [
            "The system provides all the functions I need to complete my scholarship applications.",
            "The system produces correct and appropriate results based on the academic information and documents I submit.",
            "The available functions help me search, apply for, and track scholarships effectively."
        ],
        "2. Performance Efficiency": [
            "The system responds within an acceptable amount of time when submitting forms.",
            "Scholarship catalogs, dashboards, and uploaded document previews load without unnecessary delay.",
            "The system performs satisfactorily even during peak application periods."
        ],
        "3. Compatibility": [
            "The system works properly with the laptop, smartphone, or browser (Chrome, Edge, Safari) I use.",
            "When uploading or viewing PDF and image files (COG/COR), the files render as expected across devices."
        ],
        "4. Interaction Capability": [
            "It is easy to understand what the scholarship portal is intended to do.",
            "I can learn how to apply for scholarships and track my status with minimal assistance.",
            "The menus, buttons, labels, and form instructions are clear and easy to understand.",
            "The system helps prevent mistakes or provides helpful warnings when I omit required fields.",
            "The interface is clean, organized, readable, and comfortable to use."
        ],
        "5. Reliability": [
            "The scholarship portal works consistently without unexpected crashes during submission.",
            "The portal is available and accessible whenever I need to check my application status.",
            "The system handles weak internet connections without corrupting or losing my entered application data.",
            "If an interruption occurs, the auto-save feature restores my application draft reliably."
        ],
        "6. Security": [
            "The system allows access to my personal and academic records only to authorized OSA evaluators.",
            "I feel that my submitted grades, income details, and personal records are adequately protected under R.A. 10173.",
            "The system appropriately verifies my identity via secure institutional student login.",
            "My application submissions and document updates are recorded with clear timestamps and receipts."
        ],
        "7. Flexibility": [
            "The system can support different scholarship types (institutional, government, private/corporate) seamlessly.",
            "The system remains fully usable when switching between desktop screens and mobile phone displays.",
            "The system accommodates updates to my student profile and contact details without disrupting my active applications."
        ],
        "8. Safety": [
            "The system provides appropriate warnings or confirmation dialogs before I submit or cancel an application.",
            "The system helps reduce the risk of accidental duplicate submissions or irreversible mistakes during filing."
        ]
    }

    build_part2_iso_questionnaire(
        doc,
        role_title="Student Applicant Role",
        respondent_role_box="☒ Student Applicant   ☐ OSA Evaluator / Administrator   ☐ IT Expert / Faculty Evaluator",
        statements_dict=student_iso_stmts
    )

    out_file = os.path.abspath("docs/UAT_Test_Script_Student_Role.docx")
    doc.save(out_file)
    print(f"Generated Student Document: {out_file} ({os.path.getsize(out_file)} bytes)")


# ==============================================================================
# 2. BUILD STAFF EVALUATOR ROLE DOC
# ==============================================================================
def create_staff_document():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.55)
        s.bottom_margin = Inches(0.55)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    build_part1_header(
        doc,
        role_title="OSA SCHOLARSHIP EVALUATOR / STAFF ROLE",
        role_subtitle="Staff Evaluation, AI Verification & Decisioning Validation"
    )

    build_part1_metadata_table(
        doc,
        role_default="OSA Scholarship Evaluator / Staff Officer"
    )

    staff_scenarios = [
        {
            "module": "Authentication & MFA Security",
            "task": "Staff logs in with institutional credentials, inputs 6-digit OTP, and registers trusted device token.",
            "expected": "Authentication succeeds; session encryption and role redirection route user directly to Staff Review Queue.",
        },
        {
            "module": "Application Queue Triage",
            "task": "Staff filters review queue by scholarship program, term, status (Pending/Returned), and sorts by GWA.",
            "expected": "Queue filters rapidly with real-time record count; priority badges highlight overdue or flagged applications.",
        },
        {
            "module": "Applicant Dossier Assessment",
            "task": "Staff opens application dossier to inspect student academic record, income bracket, and custom responses.",
            "expected": "Clean two-column layout renders full student data; automated eligibility badge verifies GWA compliance.",
        },
        {
            "module": "Interactive Canvas Viewer",
            "task": "Staff inspects Certificate of Grades (COG) using zoom (up to 400%), pan, rotate, and contrast inversion filters.",
            "expected": "Canvas renders smoothly with zero lag; high-contrast filter clearly exposes registrar seal details and eraser marks.",
        },
        {
            "module": "AI Fraud Score & EXIF Fingerprint",
            "task": "Staff evaluates AI Forensic score (0-100%), 3-tier risk badge, and EXIF software metadata analysis.",
            "expected": "Score and risk tier badge display accurately; benign scanner noise is properly differentiated from heavy edits.",
        },
        {
            "module": "Grad-CAM Heatmap & ELA Overlay",
            "task": "Staff toggles Grad-CAM saliency heatmap and ELA overlay on COG canvas, adjusting opacity from 0% to 100%.",
            "expected": "Overlay aligns with document coordinates; localized pixel anomalies and altered grade numbers glow prominently.",
        },
        {
            "module": "Decisioning & Fast-Triage Remarks",
            "task": "Staff renders decision (Approve/Return/Reject), selects Fast-Triage preset remarks, and confirms action.",
            "expected": "Status updates immediately in database; mandatory remark enforced on returns; automated notification triggered.",
        },
        {
            "module": "Internal Notes & Audit Log",
            "task": "Staff saves confidential internal notes and verifies that evaluation history is recorded in audit timeline.",
            "expected": "Notes remain strictly hidden from student view; immutable timeline permanently records evaluator ID and timestamp.",
        }
    ]
    build_part1_scenarios_table(doc, staff_scenarios)

    # Specialized Section B.1: Blind Review Worksheet for AI testing
    p_b1_hdr = doc.add_paragraph()
    p_b1_hdr.paragraph_format.space_before = Pt(6)
    p_b1_hdr.paragraph_format.space_after = Pt(2)
    p_b1_hdr.paragraph_format.keep_with_next = True
    r_b1 = p_b1_hdr.add_run("B.1. AI-ASSISTED VS. HUMAN-ONLY DOCUMENT REVIEW COMPARISON WORKSHEET")
    r_b1.bold = True
    r_b1.font.name = "Arial"
    r_b1.font.size = Pt(9.5)

    p_b1_txt = doc.add_paragraph()
    p_b1_txt.paragraph_format.space_before = Pt(0)
    p_b1_txt.paragraph_format.space_after = Pt(3)
    p_b1_txt.add_run(
        "Protocol: As detailed in Chapter IV (Results & Discussion), evaluators inspect 10 sample COG documents in two phases:\n"
        "Phase 1: Record authenticity judgment WITHOUT viewing AI scores. | Phase 2: Reveal AI score and heatmap, then record final decision."
    ).font.size = Pt(8)

    tbl_bld = doc.add_table(rows=11, cols=7)
    tbl_bld.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_bld)

    bld_headers = ["Doc ID", "Document Description", "Phase 1 Judgment\n(No AI Assistance)", "Review\nTime", "AI Fraud\nScore", "Phase 2 Judgment\n(With AI Assistance)", "Reviewer Confidence\n(1 to 5)"]
    bld_widths = [Inches(0.6), Inches(1.8), Inches(1.3), Inches(0.6), Inches(0.7), Inches(1.3), Inches(1.1)]

    for j, bh in enumerate(bld_headers):
        c = tbl_bld.cell(0, j)
        c.width = bld_widths[j]
        set_cell_shading(c, SHD_GREEN)
        set_cell_margins(c, 50, 50, 40, 40)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(bh)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(7.5)

    sample_docs = [
        ("COG-01", "Authentic Registrar COG (BSIT)", "☐ Auth  ☐ Tampered", "___ s", "12.4%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-02", "Grade Digit Inflation (BSCE)", "☐ Auth  ☐ Tampered", "___ s", "88.7%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-03", "Authentic Dean's List (BSA)", "☐ Auth  ☐ Tampered", "___ s", "08.1%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-04", "Forged Signature & Seal", "☐ Auth  ☐ Tampered", "___ s", "92.3%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-05", "Camera Noise Scan (BSEE)", "☐ Auth  ☐ Tampered", "___ s", "24.5%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-06", "Photoshop Spliced GWA", "☐ Auth  ☐ Tampered", "___ s", "95.6%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-07", "Authentic Academic Copy (BSBio)", "☐ Auth  ☐ Tampered", "___ s", "14.2%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-08", "Altered Units/Subjects", "☐ Auth  ☐ Tampered", "___ s", "78.4%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-09", "Mobile CamScanner Auth", "☐ Auth  ☐ Tampered", "___ s", "28.0%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5"),
        ("COG-10", "Deep Tampered Header", "☐ Auth  ☐ Tampered", "___ s", "84.9%", "☐ Auth  ☐ Tampered", "☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5")
    ]

    for row_idx, data_tuple in enumerate(sample_docs, 1):
        row = tbl_bld.rows[row_idx]
        for col_idx, val in enumerate(data_tuple):
            c = row.cells[col_idx]
            c.width = bld_widths[col_idx]
            set_cell_margins(c, 40, 40, 40, 40)
            if row_idx % 2 == 1:
                set_cell_shading(c, SHD_LIGHT_ROW)
            p = c.paragraphs[0]
            if col_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(7.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    build_part1_issue_log_and_result(doc, role_label="OSA Scholarship Evaluator")

    # 26 ISO Statements contextualized for staff
    staff_iso_stmts = {
        "1. Functional Suitability": [
            "The system provides the functions I need to triage, evaluate, and decision scholarship applications.",
            "The system produces correct fraud risk indicators and document evaluation outputs based on applicant submissions.",
            "The available review functions help me complete document integrity audits effectively and without omitting required checks."
        ],
        "2. Performance Efficiency": [
            "The review queue and applicant dossier screens respond within an acceptable amount of time.",
            "High-resolution COG documents, AI Grad-CAM heatmaps, and ELA overlays load without unnecessary delay.",
            "The system performs satisfactorily during high-volume application review sessions."
        ],
        "3. Compatibility": [
            "The evaluation canvas and review tools work properly across our office workstations and modern web browsers.",
            "When viewing scanned documents, images, and PDF uploads from students, the formatting displays accurately."
        ],
        "4. Interaction Capability": [
            "It is easy to understand what the evaluation interface and forensic indicators are intended to convey.",
            "I can learn how to navigate the review queue and interpret AI fraud scores with minimal training.",
            "The buttons, canvas controls (zoom, pan, rotate, invert), labels, and triage options are clear and easy to understand.",
            "The system helps prevent mistakes by enforcing confirmation dialogs and mandatory remarks on returned applications.",
            "The review interface is well-organized, readable, and comfortable for extended evaluation sessions."
        ],
        "5. Reliability": [
            "The evaluation portal works consistently during normal administrative working hours.",
            "The system and AI verification pipeline are available and usable whenever applications require review.",
            "The system handles corrupted file uploads gracefully without crashing the evaluation queue.",
            "If an unexpected interruption occurs, all saved evaluator remarks and decisions are preserved accurately."
        ],
        "6. Security": [
            "The system allows access to applicant dossiers and internal notes only to authorized evaluation staff.",
            "Confidential student records and financial proofs are adequately protected against unauthorized viewing.",
            "The system appropriately enforces multi-factor authentication (MFA) and trusted device access.",
            "Evaluator decisions, status changes, and timestamps are immutably logged for administrative accountability."
        ],
        "7. Flexibility": [
            "The system can support evaluation across diverse scholarship programs with varied eligibility rules.",
            "The review canvas remains usable across different monitor resolutions and responsive workstation layouts.",
            "The system accommodates customized program questions and varied document requirements without friction."
        ],
        "8. Safety": [
            "The system provides clear warnings before irreversible actions such as rejecting or revoking a scholarship grant.",
            "The system helps reduce the risk of accidental status modifications through structured confirmation modals."
        ]
    }

    build_part2_iso_questionnaire(
        doc,
        role_title="OSA Scholarship Evaluator / Staff Role",
        respondent_role_box="☐ Student Applicant   ☒ OSA Evaluator / Administrator   ☐ IT Expert / Faculty Evaluator",
        statements_dict=staff_iso_stmts
    )

    out_file = os.path.abspath("docs/UAT_Test_Script_Staff_Role.docx")
    doc.save(out_file)
    print(f"Generated Staff Document: {out_file} ({os.path.getsize(out_file)} bytes)")


# ==============================================================================
# 3. BUILD ADMIN / DIRECTOR ROLE DOC
# ==============================================================================
def create_admin_document():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.55)
        s.bottom_margin = Inches(0.55)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    build_part1_header(
        doc,
        role_title="SUPER ADMINISTRATOR & OSA DIRECTOR ROLE",
        role_subtitle="Executive Governance, Program Management & Compliance Validation"
    )

    build_part1_metadata_table(
        doc,
        role_default="Director / Head, OSA & System Super Administrator"
    )

    admin_scenarios = [
        {
            "module": "Executive KPI Analytics",
            "task": "Director reviews dashboard metrics: Total Applications, Approval Rates, Grade Integrity Index, and Quota burn.",
            "expected": "Real-time metric counters and distribution charts render accurately; dynamic academic term filtering updates totals.",
        },
        {
            "module": "Program Lifecycle Governance",
            "task": "Admin creates new scholarship program, sets slot quotas, GWA minimums, deadlines, and toggles Active/Archived.",
            "expected": "Program persists in database; slot validation prevents negative integers; active grants become visible in student catalog.",
        },
        {
            "module": "Dynamic Custom Field Builder",
            "task": "Admin configures custom scholarship form fields (text, dropdown, file), tests reorder arrows and Required switch.",
            "expected": "Dedicated sub-header renders clean button cluster without overlapping switches; field reindexing works smoothly.",
        },
        {
            "module": "RBAC & Staff Governance",
            "task": "Admin creates staff evaluator accounts, assigns specific scholarship program queues, and manages permissions.",
            "expected": "Strict RBAC restricts evaluators to assigned programs; suspended staff accounts are immediately blocked from entry.",
        },
        {
            "module": "AI Pipeline & Sensitivity Config",
            "task": "Admin inspects AI microservice connectivity, sets default pipeline (V2 ELA-CNN), and calibrates fraud score threshold.",
            "expected": "Health ping confirms live service; configured threshold updates system_settings and applies to future scan jobs.",
        },
        {
            "module": "Compliance Reporting & Exports",
            "task": "Admin exports filtered scholarship masterlists and compliance summaries in CSV and official PDF formats.",
            "expected": "CSV downloads formatted for CHED/DOST portal upload; PDF generates with CLSU OSA header, seal, and signatory lines.",
        },
        {
            "module": "System Audit Trail Monitoring",
            "task": "Admin inspects immutable audit logs, filtering by user, IP address, and action type (Logins, Overrides, Exports).",
            "expected": "Logs capture complete timestamp, actor IP, action category, and payload difference; entries cannot be altered.",
        },
        {
            "module": "Soft-Deletion & Resilience",
            "task": "Admin soft-deletes a scholarship program, inspects Trash recovery tab, and restores record back to active state.",
            "expected": "Soft-delete protects relational integrity; restore brings record back cleanly; university announcement broadcasts to feed.",
        }
    ]
    build_part1_scenarios_table(doc, admin_scenarios)
    build_part1_issue_log_and_result(doc, role_label="OSA Director / Super Administrator")

    # 26 ISO Statements contextualized for admin/director
    admin_iso_stmts = {
        "1. Functional Suitability": [
            "The system provides complete executive functions to govern scholarship programs, quotas, and compliance reporting.",
            "The system produces accurate statistical aggregates, compliance masterlists, and tamper-evident audit logs.",
            "The available administrative features help the Office of Student Affairs manage the full scholarship lifecycle effectively."
        ],
        "2. Performance Efficiency": [
            "The executive analytics dashboard and administrative management screens load within acceptable timeframes.",
            "Exportable CSV reports and authenticated PDF compliance summaries generate rapidly without server timeouts.",
            "The system maintains high performance efficiency during concurrent administrative and evaluation operations."
        ],
        "3. Compatibility": [
            "The administrative portal operates reliably across executive workstations, office PCs, and modern browsers.",
            "Exported CSV and PDF compliance files integrate cleanly with institutional spreadsheets and CHED/DOST reporting portals."
        ],
        "4. Interaction Capability": [
            "It is easy to understand the high-level metrics, grade integrity indices, and administrative settings of the system.",
            "Administrative personnel can learn to configure new scholarship programs and custom fields with minimal guidance.",
            "The navigation bars, buttons, custom field reorder controls, and program toggles are clear and intuitive.",
            "The system prevents catastrophic administrative mistakes by requiring confirmation before destructive actions.",
            "The administrative dashboard is visually balanced, professional, and well-structured for strategic oversight."
        ],
        "5. Reliability": [
            "The administrative platform operates continuously and reliably without unexpected service interruptions.",
            "The system is available and accessible whenever executive reviews or statutory compliance reports are required.",
            "The system handles heavy transaction volumes without data corruption or database deadlocks.",
            "The soft-deletion and trash recovery mechanisms protect institutional records against accidental permanent data loss."
        ],
        "6. Security": [
            "The system strictly enforces role-based access control (RBAC), restricting administrative tools to authorized personnel.",
            "Sensitive student PII, financial background data, and staff notes are encrypted at rest and in transit under R.A. 10173.",
            "The system appropriately enforces multi-factor authentication (MFA) and monitors active sessions system-wide.",
            "All administrative actions, data exports, and status overrides are permanently recorded in an immutable audit trail."
        ],
        "7. Flexibility": [
            "The system accommodates diverse scholarship structures, funding models, and academic eligibility requirements.",
            "The system scales seamlessly to support institutional growth and increased scholarship applicant volumes.",
            "The dynamic custom form field builder provides complete flexibility to collect unique program requirements as needs change."
        ],
        "8. Safety": [
            "The system provides explicit safety confirmation modals before critical actions such as deleting programs or resetting data.",
            "The system architecture incorporates safeguards that eliminate single points of failure and prevent catastrophic data corruption."
        ]
    }

    build_part2_iso_questionnaire(
        doc,
        role_title="Super Administrator & OSA Director Role",
        respondent_role_box="☐ Student Applicant   ☒ OSA Evaluator / Administrator   ☐ IT Expert / Faculty Evaluator",
        statements_dict=admin_iso_stmts
    )

    out_file = os.path.abspath("docs/UAT_Test_Script_Admin_Role.docx")
    doc.save(out_file)
    print(f"Generated Admin Document: {out_file} ({os.path.getsize(out_file)} bytes)")

if __name__ == "__main__":
    print("Beginning generation of role-specific documents aligned with Client_Testing_and_ISO25010_End_User_Evaluation.docx...")
    create_student_document()
    create_staff_document()
    create_admin_document()
    print("All 3 role-specific documents generated successfully!")
