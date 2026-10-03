#!/usr/bin/env python3
"""
A.E.G.I.S. Standardized System User & Instructional Operations Manual DOCX Generator
Generates a publication-grade, institutional Microsoft Word document (.docx)
with embedded high-resolution screenshots, formatted tables, and styled callouts.
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    """Set cell background color"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding (in dxa: 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title="IMPORTANT NOTICE"):
    """Adds a stylish callout box with a colored left border"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0FDF4") # subtle emerald tint
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border styling
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="0C4E2D"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}: ")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(12, 78, 45) # CLSU Green
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure(doc, img_path, caption_text, width_inches=6.0):
    """Embeds an image centered with a standardized figure caption"""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_text)
        run_cap.italic = True
        run_cap.font.name = "Arial"
        run_cap.font.size = Pt(9)
        run_cap.font.color.rgb = RGBColor(100, 116, 139) # text-muted
    else:
        p_err = doc.add_paragraph(f"[Image placeholder: {caption_text}]")
        p_err.italic = True

def build_manual():
    print("Initializing A.E.G.I.S. Operations Manual DOCX Generation...")
    doc = Document()
    
    # Configure 1-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        
    # Color palette
    CLSU_GREEN = RGBColor(12, 78, 45)
    CLSU_GOLD = RGBColor(242, 169, 0)
    DARK_SLATE = RGBColor(15, 23, 42)
    MUTED_GRAY = RGBColor(100, 116, 139)
    
    # ── COVER / TITLE HEADER ──────────────────────────────────────────────────
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(12)
    p_inst.paragraph_format.space_after = Pt(2)
    run_inst = p_inst.add_run("CENTRAL LUZON STATE UNIVERSITY\nOFFICE OF STUDENT AFFAIRS (OSA)")
    run_inst.bold = True
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(11)
    run_inst.font.color.rgb = CLSU_GREEN
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("A.E.G.I.S.")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(28)
    run_title.font.color.rgb = CLSU_GREEN
    
    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subtitle.paragraph_format.space_after = Pt(6)
    run_sub = p_subtitle.add_run("Academic Evaluation & Grant Integrity System")
    run_sub.bold = True
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = DARK_SLATE
    
    p_doc = doc.add_paragraph()
    p_doc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_doc.paragraph_format.space_after = Pt(18)
    run_doc = p_doc.add_run("STANDARDIZED SYSTEM USER & INSTRUCTIONAL OPERATIONS MANUAL\nDocument Reference: CLSU-OSA-AEGIS-UIM-2026-V1 │ Production Release v1.0.0")
    run_doc.font.name = "Arial"
    run_doc.font.size = Pt(10)
    run_doc.font.color.rgb = MUTED_GRAY

    # Document Control Table
    table_ctrl = doc.add_table(rows=5, cols=2)
    table_ctrl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_ctrl.autofit = False
    
    metadata = [
        ("System Classification", "Enterprise Web Portal & Deep Learning AI Forensics Microservice"),
        ("Statutory Compliance", "R.A. 10173 (Data Privacy Act of 2012) & R.A. 11032 (Ease of Doing Business Act)"),
        ("Security Standards", "OWASP Top 10 Hardened, AES-256-CBC Column Encryption, Blind Indexing"),
        ("Target Users", "CLSU Enrolled Students, OSA Evaluators (Staff), OSA Director (SuperAdmin), DevOps Engineers"),
        ("Production Host & DB", "Docker Multi-Stage Runtime (Nginx + PHP 8.4-FPM) / PostgreSQL (Supabase)")
    ]
    
    for idx, (label, val) in enumerate(metadata):
        c0 = table_ctrl.cell(idx, 0)
        c1 = table_ctrl.cell(idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Arial"
        r0.font.size = Pt(9)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)

    doc.add_page_break()

    # ── SECTION 1: ARCHITECTURE & ACCESS PREREQUISITES ─────────────────────────
    h1 = doc.add_heading("1. System Overview & Access Architecture", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(8)
    
    doc.add_paragraph(
        "A.E.G.I.S. is an automated scholarship management and integrity auditing portal developed "
        "for the Central Luzon State University (CLSU) Office of Student Affairs (OSA). It automates "
        "application intake, verifies General Weighted Averages (GWA) via Optical Character Recognition (OCR), "
        "and detects digital document tampering using Error Level Analysis (ELA) combined with ResNet-50 "
        "Convolutional Neural Networks."
    )
    
    add_callout(
        doc,
        "All scholarship applicants and administrative reviewers must access the portal using modern "
        "web standards (Google Chrome v110+, Mozilla Firefox v115+, or Apple Safari v16+) over secure HTTPS connections. "
        "Student file uploads are restricted to 10.0 MB per document in PDF, JPG, or PNG formats.",
        "ACCESS PREREQUISITE"
    )

    # ── SECTION 2: STUDENT PORTAL OPERATIONS MANUAL ───────────────────────────
    h2 = doc.add_heading("2. Module 1: Student Portal Operations Manual", level=1)
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(8)

    doc.add_heading("2.1 Registration & Institutional Email Verification", level=2)
    doc.add_paragraph(
        "1. Open your web browser and navigate to https://your-domain.clsu.edu.ph/register.\n"
        "2. Provide your full legal name, official CLSU institutional email address (@clsu.edu.ph), "
        "and configure a strong password (minimum 8 characters with mixed case, numbers, and symbols).\n"
        "3. Click 'Create Student Account'.\n"
        "4. A cryptographically signed activation token will be delivered to your university Gmail inbox. "
        "Click 'Verify Email Address' in the confirmation email to activate your account."
    )

    doc.add_heading("2.2 Two-Factor Authentication & Device Trust", level=2)
    doc.add_paragraph(
        "A.E.G.I.S. enforces strict two-factor authentication (6-digit OTP) on all logins to safeguard student academic records."
    )
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_01_Login_Page.png",
        "Figure 1: Institutional Login & Two-Factor Authentication (MFA) Portal",
        5.8
    )
    doc.add_paragraph(
        "Operational Instructions for Figure 1:\n"
        "• [1] Login Input: Enter your institutional email address and password in the credentials panel.\n"
        "• [2] Two-Factor Verification Prompt: Enter the 6-digit OTP code received in your inbox (valid for 10 minutes).\n"
        "• [3] 30-Day Device Trust: Check 'Remember this trusted device for 30 days' when logging in on private hardware. "
        "The system generates a cryptographically signed cookie bound to your browser User-Agent hash, bypassing OTP on subsequent logins."
    )

    doc.add_heading("2.3 Student Profile Completion & Data Privacy Act Compliance", level=2)
    doc.add_paragraph(
        "Before filing an application, students must complete their permanent institutional profile:\n"
        "• CLSU Student ID Number: Enter in standardized XX-XXXX format (e.g., 23-1234). The system computes an automated "
        "SHA-256 blind index hash to prevent duplicate accounts across students while encrypting the plaintext ID at rest using AES-256-CBC.\n"
        "• Degree Program Catalog: Select your college to cascade all official CHED degree programs dynamically.\n"
        "• Permanent Home Address: Select Province, City/Municipality, and Barangay using the live Philippine Standard Geographic Code (PSGC) API.\n"
        "• Guardian & Contact Information: Provide an 11-digit Philippine mobile contact number (09XXXXXXXXX)."
    )

    doc.add_heading("2.4 Student Dashboard & Lifecycle Pizza Tracker", level=2)
    doc.add_paragraph(
        "Upon authentication, the student is greeted by the visual lifecycle Pizza Tracker, displaying live application status."
    )
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_13_Student_Dashboard.png",
        "Figure 2: Student Dashboard & Real-Time Lifecycle Pizza Tracker",
        5.8
    )
    doc.add_paragraph(
        "Operational Instructions for Figure 2:\n"
        "• [1] Pizza Tracker Progression: Visually tracks status through 5 phases: Submitted → Under Review → Integrity Scanned → Approved → Disbursed.\n"
        "• [2] Active Grant Card: Shows current grant name, academic term, stipend grant amount, and renewal count.\n"
        "• [3] Evaluator Remarks: View feedback, requests for clearer document scans, or office-hour queuing notices."
    )

    doc.add_heading("2.5 Submitting a Scholarship Application (3-Step Stepper)", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_14_Student_Apply_Form.png",
        "Figure 3: Student 3-Step Interactive Application Stepper",
        5.8
    )
    doc.add_paragraph(
        "Step-by-Step Submission Procedure (Figure 3):\n"
        "1. Step 1 (Program Selection): Select your desired scholarship program from the active institutional catalog [1].\n"
        "2. Step 2 (Academic Declarations): Enter your declared General Weighted Average (GWA) from the previous semester (e.g., 1.45) "
        "and complete any program-specific criteria [2]. Note: In the Philippine grading scale, 1.00 represents highest honors and 3.00 is passing.\n"
        "3. Step 3 (Document Upload): Attach an uncropped, high-resolution scan of your official Certificate of Grades (COG) [3].\n"
        "4. Click 'Submit Application'. If submitted after 5:00 PM PHT on weekdays or during weekends, the application is queued for evaluation on the next business morning."
    )

    doc.add_heading("2.6 Document Re-uploads & Official PDF Download", level=2)
    doc.add_paragraph(
        "• Re-uploading Documents: If an evaluator requests an updated scan due to blur or lighting issues, an 'Upload Clear Copy' button appears on your dashboard.\n"
        "• Downloading Approval Form: Once awarded, click 'Download Official Approved Application Form (PDF)' to generate the official signed document bearing the university seal and stipend claim guidelines."
    )

    doc.add_page_break()

    # ── SECTION 3: OSA EVALUATOR (STAFF) OPERATIONS MANUAL ─────────────────────
    h3 = doc.add_heading("3. Module 2: OSA Evaluator (Staff) Operations Manual", level=1)
    h3.paragraph_format.space_before = Pt(18)
    h3.paragraph_format.space_after = Pt(8)

    doc.add_heading("3.1 Application Queue Triage & Priority Sorting", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_07_Admin_Application_Queue.png",
        "Figure 4: OSA Evaluator Application Queue & Triage Console",
        5.8
    )
    doc.add_paragraph(
        "Operational Instructions for Figure 4:\n"
        "• [1] Metrics Summary Bar: View real-time pending submissions, under-review applications, and flagged tampering cases.\n"
        "• [2] Dynamic Filter Panel: Filter by program, term, evaluation status, or workload assignment ('Assigned to Me' vs 'Unassigned').\n"
        "• [3] Priority Risk Sorting: Choose 'Priority Risk' to elevate applications with AI fraud scores ≥ 70.0% to the top of the queue.\n"
        "• [4] Action Button: Click 'Review Applicant' to enter the forensic evaluation studio."
    )

    doc.add_heading("3.2 The 4-Pillar Forensic Review Studio", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_09_Application_Review_Detail.png",
        "Figure 5: Forensic Review Studio & ELA Heatmap Analysis",
        5.8
    )
    doc.add_paragraph(
        "A.E.G.I.S. provides an Explainable Forensic Decision Framework (EFDF) structured across four transparent pillars (Figure 5):\n"
        "• Pillar 1 (Document Syntax Gate): Pre-filters invalid files (memes, selfies, unrelated graphics).\n"
        "• Pillar 2 (OCR Grade Consistency — 35% Weight): Compares student-declared GWA against Tesseract-extracted values.\n"
        "• Pillar 3 (Compression & ELA Forensics — 50% Weight): Error Level Analysis (ELA) and ResNet-50 CNN detect spliced or altered grade digits, highlighted in thermal magenta/yellow heatmaps [2].\n"
        "• Pillar 4 (Metadata Provenance — 15% Weight): Native EXIF analysis inspects image headers for editing software (Photoshop, Photopea)."
    )

    doc.add_heading("3.3 GWA 4-Tier Discrepancy Calibration", level=2)
    
    # Discrepancy Table
    table_disc = doc.add_table(rows=5, cols=3)
    table_disc.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_disc.autofit = False
    
    disc_data = [
        ("Tier", "Discrepancy Condition", "System Action & Decision Guidance"),
        ("Tier 1", "Verified Match (Diff ≤ 0.01)", "GWA verified; 0% penalty applied. Normal approval flow."),
        ("Tier 2", "Minor OCR Variance (Diff ≤ 0.05)", "Single-digit OCR misread or rounding. Advisory review flag; student is NOT accused of fraud."),
        ("Tier 3", "Grade Inflation (Declared < Extracted)", "Student declared a better numeric grade than transcript reflects. 99.0% fraud score; flagged as Tampered."),
        ("Tier 4", "Inverse Variance (Declared > Extracted)", "Student declared a worse grade than earned (data-entry mistake). Advisory flag for staff correction.")
    ]
    
    for r_idx, row in enumerate(disc_data):
        for c_idx, val in enumerate(row):
            cell = table_disc.cell(r_idx, c_idx)
            set_cell_margins(cell, 80, 80, 100, 100)
            if r_idx == 0:
                set_cell_background(cell, "0C4E2D")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(val)
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.add_heading("3.4 Announcement Board Management", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_11_Announcements.png",
        "Figure 6: Institutional Announcement Management Console",
        5.8
    )
    doc.add_paragraph(
        "Staff evaluators can broadcast deadline extensions, missing document notices, and stipend schedules to all student dashboards (Figure 6)."
    )

    doc.add_page_break()

    # ── SECTION 4: OSA DIRECTOR (SUPERADMIN) OPERATIONS MANUAL ─────────────────
    h4 = doc.add_heading("4. Module 3: OSA Director (SuperAdmin) Operations Manual", level=1)
    h4.paragraph_format.space_before = Pt(18)
    h4.paragraph_format.space_after = Pt(8)

    doc.add_heading("4.1 Scholarship Program & Quota Management", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_02_Scholarship_Programs.png",
        "Figure 7: Scholarship Program Configuration & Quotas",
        5.8
    )
    doc.add_paragraph(
        "The OSA Director manages grant policies (Figure 7):\n"
        "• Create Programs [1]: Define funding sources, minimum GWA cutoffs, maximum renewal counts, and stipend grant amounts.\n"
        "• Quota Caps: Set maximum beneficiaries per semester or leave uncapped.\n"
        "• Dynamic Form Schema: Attach required document checklists (e.g., Certificate of Indigency, Solo Parent Certificate)."
    )

    doc.add_heading("4.2 Executive Analytics & ISO/IEC 25010 Quality Metrics", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_03_Analytics_Dashboard.png",
        "Figure 8: Executive Analytics Dashboard & ISO/IEC 25010 Quality Radar",
        5.8
    )
    doc.add_paragraph(
        "Real-time executive intelligence features (Figure 8):\n"
        "• [1] Grade Integrity Index: Mean integrity score across all awarded scholarships (100% - Avg Fraud Score).\n"
        "• [2] Turnaround Velocity: Average evaluation cycle time in days.\n"
        "• [3] College Distribution: Grant allocation breakdown across CLSU's 9 colleges.\n"
        "• [4] GWA Distribution Histogram: Density curve of applicant academic ratings.\n"
        "• [5] ISO/IEC 25010 Evaluator Ratings: Live radar of user acceptance across Functional Suitability, Usability, Reliability, and Security."
    )

    doc.add_heading("4.3 Staff Delegation & Workload Partitioning", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_06_Staff_Management.png",
        "Figure 9: Staff Delegation & Role Management Console",
        5.8
    )
    doc.add_paragraph(
        "Directors invite new evaluators via secure single-use email tokens [1] and assign specific scholarship programs [2] "
        "to partition application volume evenly across personnel (Figure 9)."
    )

    doc.add_heading("4.4 Dynamic System Configuration & AI Sensitivity", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_04_System_Settings.png",
        "Figure 10: Dynamic System Settings & Forensics Calibration",
        5.8
    )
    doc.add_paragraph(
        "Global configuration controls (Figure 10):\n"
        "• AI Fraud Threshold: Default set to 70.0% to balance mobile camera compression artifacts with genuine forgery detection.\n"
        "• GWA Tolerance: Set mathematical discrepancy margin (Default: 0.01).\n"
        "• MFA Policy: Enforce for all accounts, staff only, or optional.\n"
        "• Universal Device Revocation: Immediately terminates all 30-day remembered browser cookies across the university."
    )

    doc.add_heading("4.5 7-Tier Compliance Audits & Official Log Exports", level=2)
    add_figure(
        doc,
        "thesis_figures/screenshots/Figure_05_Audit_Logs.png",
        "Figure 11: System Audit Trail & Compliance Log Center",
        5.8
    )
    doc.add_paragraph(
        "A.E.G.I.S. logs every system event across 7 dedicated security tiers: Authentication, Admin Actions, AI Scans, "
        "Evaluation Decisions, Configuration Changes, Email Deliveries, and Export Access Logs (Figure 11). "
        "All logs are exportable to official CSV and PDF formats for state university auditing."
    )

    doc.add_page_break()

    # ── SECTION 5: DEVOPS & SYSTEM ADMINISTRATION REFERENCE ───────────────────
    h5 = doc.add_heading("5. Module 4: System Administration & DevOps Reference", level=1)
    h5.paragraph_format.space_before = Pt(18)
    h5.paragraph_format.space_after = Pt(8)

    doc.add_paragraph(
        "Production Container Architecture:\n"
        "• Multi-Stage Dockerfile: Stage 1 (Node 20 asset compiler), Stage 2 (Composer 2 vendor optimizer), "
        "Stage 3 (Alpine PHP 8.4-FPM + Nginx + Supervisord runtime).\n"
        "• Process Supervisor: Supervisord monitors PHP-FPM, Nginx, and two background queue workers (php artisan queue:work --tries=3).\n"
        "• Sentry Error Monitoring: Integrated via \\Sentry\\Laravel\\Integration::handles($exceptions) in bootstrap/app.php.\n"
        "• Composite High-Frequency Database Indexes: idx_apps_status_archived, idx_apps_assigned_status, idx_apps_user_term, idx_apps_created_at."
    )

    doc.add_heading("6. Troubleshooting Matrix & Frequently Asked Questions (FAQ)", level=1)
    
    table_faq = doc.add_table(rows=4, cols=3)
    table_faq.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_faq.autofit = False
    
    faq_data = [
        ("Symptom", "Probable Cause", "Remediation Procedure"),
        ("MFA 6-digit code not received in inbox.", "Email spam filter or Brevo delivery delay.", "Check spam/junk folders. Evaluators verify delivery status via /superadmin/email-logs. Whitelist noreply@clsu-aegis.ph."),
        ("Uploaded COG image rejected as 'Tampered'.", "Heavy mobile photo editing or extreme compression.", "Advise student to take an unedited, original camera photograph of the printed Certificate of Grades under natural daylight."),
        ("Application marked as 'Queued for Next Business Day'.", "Submitted outside regular OSA office hours.", "Normal operational procedure. Queue automatically opens at 8:00 AM on the following business day.")
    ]
    
    for r_idx, row in enumerate(faq_data):
        for c_idx, val in enumerate(row):
            cell = table_faq.cell(r_idx, c_idx)
            set_cell_margins(cell, 80, 80, 100, 100)
            if r_idx == 0:
                set_cell_background(cell, "0C4E2D")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(val)
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(8.5)

    # Save to target destination
    out_dir = "docs"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "AEGIS_STANDARDIZED_USER_AND_INSTRUCTIONAL_MANUAL.docx")
    doc.save(out_path)
    print(f"Successfully generated operations manual DOCX: {out_path}")
    
    # Also save a copy to the root for immediate access
    root_path = "AEGIS_USER_AND_INSTRUCTIONAL_MANUAL.docx"
    doc.save(root_path)
    print(f"Successfully saved root copy: {root_path}")

if __name__ == "__main__":
    build_manual()
