"""
Convert A.E.G.I.S. Thesis from Markdown to DOCX
Professional academic formatting with proper styles
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re
import os

print("A.E.G.I.S. Thesis Markdown to DOCX Converter")
print("=" * 60)

# Create new document
doc = Document()

# Set up document margins (1 inch all sides)
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Define styles
def setup_styles(doc):
    """Setup custom styles for academic thesis"""

    styles = doc.styles

    # Normal style
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    # Heading 1 style
    heading1_style = styles['Heading 1']
    heading1_font = heading1_style.font
    heading1_font.name = 'Times New Roman'
    heading1_font.size = Pt(16)
    heading1_font.bold = True
    heading1_style.paragraph_format.space_before = Pt(12)
    heading1_style.paragraph_format.space_after = Pt(6)

    # Heading 2 style
    heading2_style = styles['Heading 2']
    heading2_font = heading2_style.font
    heading2_font.name = 'Times New Roman'
    heading2_font.size = Pt(14)
    heading2_font.bold = True
    heading2_style.paragraph_format.space_before = Pt(10)
    heading2_style.paragraph_format.space_after = Pt(6)

    # Heading 3 style
    heading3_style = styles['Heading 3']
    heading3_font = heading3_style.font
    heading3_font.name = 'Times New Roman'
    heading3_font.size = Pt(12)
    heading3_font.bold = True
    heading3_style.paragraph_format.space_before = Pt(6)
    heading3_style.paragraph_format.space_after = Pt(6)

setup_styles(doc)

# Read thesis markdown file
print("\n[1/3] Reading thesis markdown file...")
with open('AEGIS_COMPLETE_CAPSTONE2_THESIS.md', 'r', encoding='utf-8') as f:
    content = f.read()

print("[OK] Thesis file loaded successfully!")

# Parse and convert markdown to DOCX
print("\n[2/3] Converting markdown to DOCX with academic formatting...")

lines = content.split('\n')
in_table = False
table_lines = []

for line in lines:
    # Skip empty lines at document start
    if not line.strip():
        if len(doc.paragraphs) > 0 or len(doc.tables) > 0:
            doc.add_paragraph()
        continue

    # Heading 1 (# )
    if line.startswith('# ') and not line.startswith('## '):
        heading = line.replace('# ', '').strip()
        p = doc.add_heading(heading, level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Heading 2 (## )
    elif line.startswith('## ') and not line.startswith('### '):
        heading = line.replace('## ', '').strip()
        p = doc.add_heading(heading, level=2)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Heading 3 (### )
    elif line.startswith('### ') and not line.startswith('#### '):
        heading = line.replace('### ', '').strip()
        p = doc.add_heading(heading, level=3)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Heading 4 (#### )
    elif line.startswith('#### '):
        heading = line.replace('#### ', '').strip()
        p = doc.add_paragraph(heading)
        run = p.runs[0]
        run.bold = True
        run.font.size = Pt(12)

    # Horizontal rule (---)
    elif line.strip() == '---':
        doc.add_page_break()

    # Table detection (starts with |)
    elif line.strip().startswith('|'):
        if not in_table:
            in_table = True
            table_lines = []
        table_lines.append(line)

    # End of table
    elif in_table and not line.strip().startswith('|'):
        # Process collected table
        if len(table_lines) > 2:  # At least header + separator + 1 row
            # Parse table
            rows = []
            for tline in table_lines:
                if '|--' not in tline:  # Skip separator line
                    cells = [cell.strip() for cell in tline.split('|')[1:-1]]
                    if cells:
                        rows.append(cells)

            if rows:
                # Create table
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = 'Light Grid Accent 1'

                # Fill table
                for i, row_data in enumerate(rows):
                    row = table.rows[i]
                    for j, cell_data in enumerate(row_data):
                        cell = row.cells[j]
                        cell.text = cell_data
                        # Bold header row
                        if i == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True
                        # Set font
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.name = 'Times New Roman'
                                run.font.size = Pt(11)

        in_table = False
        table_lines = []

        # Now process current line as normal
        if line.strip():
            doc.add_paragraph(line)

    # Bold text (**text**)
    elif '**' in line:
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part.replace('**', ''))
                run.bold = True
            else:
                p.add_run(part)

    # Bullet points (- )
    elif line.strip().startswith('- '):
        text = line.strip()[2:]
        # Handle bold within bullet points
        if '**' in text:
            p = doc.add_paragraph(style='List Bullet')
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part.replace('**', ''))
                    run.bold = True
                else:
                    p.add_run(part)
        else:
            doc.add_paragraph(text, style='List Bullet')

    # Numbered lists (1. )
    elif re.match(r'^\d+\.\s', line.strip()):
        text = re.sub(r'^\d+\.\s', '', line.strip())
        doc.add_paragraph(text, style='List Number')

    # Regular paragraph
    else:
        doc.add_paragraph(line)

print("[OK] Markdown converted to DOCX!")

# Save document
output_file = 'AEGIS_COMPLETE_CAPSTONE2_THESIS.docx'
print(f"\n[3/3] Saving document as {output_file}...")
doc.save(output_file)

print("[OK] Document saved successfully!")

# Print summary
print("\n" + "=" * 60)
print("CONVERSION COMPLETE!")
print("=" * 60)
print(f"\nOutput File: {output_file}")
print(f"Location: {os.path.abspath(output_file)}")
print(f"\nDocument Statistics:")
print(f"  - Paragraphs: {len(doc.paragraphs)}")
print(f"  - Tables: {len(doc.tables)}")
print(f"  - Sections: {len(doc.sections)}")

print("\nFormatting Applied:")
print("  - Font: Times New Roman 12pt")
print("  - Line Spacing: 1.5")
print("  - Margins: 1 inch (all sides)")
print("  - Headings: Hierarchical (H1, H2, H3)")
print("  - Tables: Professional grid style")

print("\nNext Steps:")
print("  1. Open in Microsoft Word")
print("  2. Insert charts from thesis_figures/ folder")
print("  3. Add page numbers (Insert > Page Number)")
print("  4. Generate Table of Contents (References > Table of Contents)")
print("  5. Update TOC page numbers")
print("  6. Final proofread and save as PDF")

print("\n" + "=" * 60)
print("Ready for submission!")
print("=" * 60)
